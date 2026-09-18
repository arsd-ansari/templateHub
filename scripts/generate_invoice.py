"""
Generates a clean, general-purpose Invoice Template (.xlsx) for any business.
Line amounts, subtotal, tax and total calculate automatically.

Run:  .venv-tools/bin/python scripts/generate_invoice.py
Output: public/files/invoice-template.xlsx (shipped with the site)
"""

from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

BRAND = "0F766E"
BRAND_LIGHT = "D7EEEB"
GREY = "F1F5F9"
WHITE = "FFFFFF"
DARK = "16181D"
money = '#,##0.00'
pct = '0%'

thin = Side(style="thin", color="C9D2DC")
box = Border(left=thin, right=thin, top=thin, bottom=thin)

wb = Workbook()
ws = wb.active
ws.title = "Invoice"
ws.sheet_view.showGridLines = False
for i, w in enumerate([7, 40, 10, 15, 16], start=1):   # A..E
    ws.column_dimensions[get_column_letter(i)].width = w


def style(ref, *, value=None, bold=False, size=11, color=DARK, fill=None,
          align="left", border=False, wrap=False, number_format=None, italic=False):
    c = ws[ref]
    if value is not None:
        c.value = value
    c.font = Font(name="Calibri", bold=bold, size=size, color=color, italic=italic)
    c.alignment = Alignment(horizontal=align, vertical="center", wrap_text=wrap)
    if fill:
        c.fill = PatternFill("solid", fgColor=fill)
    if border:
        c.border = box
    if number_format:
        c.number_format = number_format
    return c


# Title
ws.merge_cells("A1:E1")
style("A1", value="INVOICE", bold=True, size=24, color=WHITE, fill=BRAND, align="center")
ws.row_dimensions[1].height = 36

# From (seller)
style("A3", value="From", bold=True, color=BRAND)
ws.merge_cells("A4:B4"); style("A4", value="[Your Company Name]", bold=True, size=13)
ws.merge_cells("A5:B5"); style("A5", value="[Street Address, City, State - PIN]")
ws.merge_cells("A6:B6"); style("A6", value="[Email]  |  [Phone]")

# Invoice meta (right)
meta = [("Invoice No.", "[INV-001]"), ("Invoice Date", "[DD-MM-YYYY]"), ("Due Date", "[DD-MM-YYYY]")]
r = 3
for label, val in meta:
    style(f"D{r}", value=label, bold=True, fill=GREY, border=True)
    style(f"E{r}", value=val, border=True, align="right")
    r += 1

# Bill To
ws.merge_cells("A8:E8")
style("A8", value="BILL TO", bold=True, color=WHITE, fill=BRAND, border=True)
for row, key in ((9, "Client Name"), (10, "Address"), (11, "Email / Phone")):
    ws.merge_cells(f"A{row}:E{row}")
    style(f"A{row}", value=f"{key}: [ ]", border=True)

# Items header
H = 13
for i, h in enumerate(["S.No", "Description", "Qty", "Rate", "Amount"], start=1):
    style(f"{get_column_letter(i)}{H}", value=h, bold=True, color=WHITE, fill=BRAND,
          align="center", border=True)
ws.row_dimensions[H].height = 24

FIRST = H + 1
N = 8
LAST = FIRST + N - 1
samples = [
    (1, "Website design - landing page", 1, 15000),
    (2, "Logo design", 1, 5000),
    (3, "Hosting setup (one-time)", 1, 2000),
]
for idx, row in enumerate(range(FIRST, LAST + 1)):
    s = samples[idx] if idx < len(samples) else None
    style(f"A{row}", value=(s[0] if s else None), align="center", border=True)
    style(f"B{row}", value=(s[1] if s else None), border=True)
    style(f"C{row}", value=(s[2] if s else None), align="center", border=True)
    style(f"D{row}", value=(s[3] if s else None), align="right", border=True, number_format=money)
    style(f"E{row}", value=f'=IF(OR($C{row}="",$D{row}=""),"",ROUND($C{row}*$D{row},2))',
          align="right", border=True, number_format=money)
    if idx % 2 == 1:
        for i in range(1, 6):
            ws[f"{get_column_letter(i)}{row}"].fill = PatternFill("solid", fgColor=GREY)


def total_row(row, label, formula, *, strong=False, is_pct=False):
    ws.merge_cells(f"A{row}:C{row}")
    style(f"A{row}", value=label, bold=True, align="right",
          fill=(BRAND if strong else BRAND_LIGHT), color=(WHITE if strong else DARK), border=True)
    style(f"D{row}", border=True, fill=(BRAND if strong else None))
    style(f"E{row}", value=formula, bold=True, align="right", border=True,
          number_format=(pct if is_pct else money),
          fill=(BRAND if strong else BRAND_LIGHT), color=(WHITE if strong else DARK),
          size=(12 if strong else 11))


sub = LAST + 1
total_row(sub, "Subtotal", f"=ROUND(SUM(E{FIRST}:E{LAST}),2)")
# Tax rate is user-editable (default 0%); tax amount + total follow automatically.
ws.merge_cells(f"A{sub+1}:C{sub+1}")
style(f"A{sub+1}", value="Tax Rate", bold=True, align="right", fill=BRAND_LIGHT, border=True)
style(f"D{sub+1}", value=0, align="right", border=True, number_format=pct, fill=BRAND_LIGHT)
style(f"E{sub+1}", value=f"=ROUND(E{sub}*D{sub+1},2)", align="right", border=True, number_format=money, fill=BRAND_LIGHT)
total_row(sub + 2, "TOTAL", f"=E{sub}+E{sub+1}", strong=True)

# Notes
n = sub + 4
style(f"A{n}", value="Notes / Payment Terms", bold=True, color=BRAND)
for i, line in enumerate([
    "Payment due within 15 days. Bank: [ ]  A/C: [ ]  IFSC: [ ]  UPI: [ ]",
    "Thank you for your business!",
]):
    ws.merge_cells(f"A{n+1+i}:E{n+1+i}")
    style(f"A{n+1+i}", value=line, italic=(i == 1))

ws.freeze_panes = f"A{FIRST}"

out = Path("public/files/invoice-template.xlsx")
out.parent.mkdir(parents=True, exist_ok=True)
wb.save(out)
print(f"Saved: {out.resolve()}")

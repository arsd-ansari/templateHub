"""
Generates a Profit & Loss (Income) Statement template (.xlsx) for small
businesses. Revenue, COGS, gross profit, expenses and net profit all total
automatically.

Run:  .venv-tools/bin/python scripts/generate_profit_and_loss.py
Output: public/files/profit-and-loss-statement.xlsx (shipped with the site)
"""

from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

BRAND = "0F766E"
BRAND_LIGHT = "D7EEEB"
GREY = "F1F5F9"
WHITE = "FFFFFF"
DARK = "16181D"
money = '#,##0.00'

thin = Side(style="thin", color="C9D2DC")
box = Border(left=thin, right=thin, top=thin, bottom=thin)

wb = Workbook()
ws = wb.active
ws.title = "Profit & Loss"
ws.sheet_view.showGridLines = False
ws.column_dimensions["A"].width = 42
ws.column_dimensions["B"].width = 20


def style(ref, *, value=None, bold=False, size=11, color=DARK, fill=None,
          align="left", border=False, number_format=None):
    c = ws[ref]
    if value is not None:
        c.value = value
    c.font = Font(name="Calibri", bold=bold, size=size, color=color)
    c.alignment = Alignment(horizontal=align, vertical="center")
    if fill:
        c.fill = PatternFill("solid", fgColor=fill)
    if border:
        c.border = box
    if number_format:
        c.number_format = number_format
    return c


def header(row, text):
    style(f"A{row}", value=text, bold=True, color=WHITE, fill=BRAND, border=True)
    style(f"B{row}", fill=BRAND, border=True)


def line(row, label, val):
    style(f"A{row}", value=label, border=True)
    style(f"B{row}", value=val, border=True, align="right", number_format=money)


def total(row, label, formula, *, strong=False):
    fill = BRAND if strong else BRAND_LIGHT
    col = WHITE if strong else DARK
    style(f"A{row}", value=label, bold=True, fill=fill, color=col, border=True,
          size=(12 if strong else 11))
    style(f"B{row}", value=formula, bold=True, fill=fill, color=col, border=True,
          align="right", number_format=money, size=(12 if strong else 11))


# Title + meta
ws.merge_cells("A1:B1")
style("A1", value="PROFIT & LOSS STATEMENT", bold=True, size=20, color=WHITE, fill=BRAND, align="center")
ws.row_dimensions[1].height = 32
ws.merge_cells("A2:B2")
style("A2", value="Business: [Your Company Name]      Period: [e.g. FY 2026-27]", color="667085")

# Revenue
header(4, "REVENUE")
line(5, "Sales revenue", 500000)
line(6, "Service revenue", 150000)
line(7, "Other income", 20000)
line(8, "[Add revenue line]", None)
line(9, "[Add revenue line]", None)
total(10, "Total Revenue", "=SUM(B5:B9)")

# Cost of goods sold
header(12, "COST OF GOODS SOLD (COGS)")
line(13, "Purchases / materials", 220000)
line(14, "Direct labour", 80000)
line(15, "[Add COGS line]", None)
line(16, "[Add COGS line]", None)
total(17, "Total COGS", "=SUM(B13:B16)")

# Gross profit
total(19, "GROSS PROFIT", "=B10-B17")

# Operating expenses
header(21, "OPERATING EXPENSES")
line(22, "Rent", 60000)
line(23, "Salaries & wages", 120000)
line(24, "Utilities", 15000)
line(25, "Marketing", 25000)
line(26, "Software & subscriptions", 12000)
line(27, "Other expenses", 10000)
line(28, "[Add expense line]", None)
line(29, "[Add expense line]", None)
total(30, "Total Operating Expenses", "=SUM(B22:B29)")

# Net profit
total(32, "NET PROFIT", "=B19-B30", strong=True)

out = Path("public/files/profit-and-loss-statement.xlsx")
out.parent.mkdir(parents=True, exist_ok=True)
wb.save(out)
print(f"Saved: {out.resolve()}")

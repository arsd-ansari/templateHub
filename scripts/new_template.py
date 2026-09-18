"""
new_template.py — scaffolder for adding a new template to ExcelHub.

Run:  .venv-tools/bin/python scripts/new_template.py

It asks you a few questions, then:
  1. Creates a starter generator at  scripts/generate_<slug>.py  (edit it to
     design your real columns/formulas, then run it to produce the .xlsx).
  2. Prints the exact block to paste into  prisma/seed.ts  (already filled with
     your title, slug, category and file path).

It does NOT write content or design your spreadsheet for you — that's the part
that needs a brain (yours, or an AI you direct with the prompts in PLAYBOOK.md).
See PLAYBOOK.md for the full step-by-step.
"""

import re
from pathlib import Path

CATEGORIES = [
    "accounting-templates", "finance-templates", "gst-templates",
    "payroll-templates", "hr-templates", "inventory-templates",
    "sales-templates", "business-templates", "education-templates",
]


def slugify(text):
    text = text.lower().strip()
    text = re.sub(r"[^a-z0-9]+", "-", text)
    return text.strip("-")


GENERATOR_SKELETON = '''"""
Generator for the "__TITLE__" template.
Edit the columns, sample rows and formulas below, then run:
    .venv-tools/bin/python scripts/generate___SLUGUNDER__.py
Output: public/files/__SLUG__.xlsx
"""

from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

BRAND = "0F766E"; GREY = "F1F5F9"; WHITE = "FFFFFF"; DARK = "16181D"
money = "#,##0.00"
thin = Side(style="thin", color="C9D2DC")
box = Border(left=thin, right=thin, top=thin, bottom=thin)

wb = Workbook()
ws = wb.active
ws.title = "Sheet1"
ws.sheet_view.showGridLines = False
# TODO: set your column widths
for i, w in enumerate([8, 36, 12, 14, 16], start=1):
    ws.column_dimensions[get_column_letter(i)].width = w


def style(ref, *, value=None, bold=False, size=11, color=DARK, fill=None,
          align="left", border=False, number_format=None):
    c = ws[ref]
    if value is not None:
        c.value = value
    c.font = Font(name="Calibri", bold=bold, size=size, color=color)
    c.alignment = Alignment(horizontal=align, vertical="center")
    if fill:
        c.fill = PatternFill("solid", fgColor=fill)
    if border:
        c.border = box
    if number_format:
        c.number_format = number_format
    return c


# Title
ws.merge_cells("A1:E1")
style("A1", value="__TITLE__".upper(), bold=True, size=20, color=WHITE, fill=BRAND, align="center")
ws.row_dimensions[1].height = 32

# TODO: design your real table. This is just a starter example.
headers = ["S.No", "Description", "Qty", "Rate", "Amount"]
for i, h in enumerate(headers, start=1):
    style(f"{get_column_letter(i)}3", value=h, bold=True, color=WHITE, fill=BRAND,
          align="center", border=True)

FIRST = 4
LAST = 11
for row in range(FIRST, LAST + 1):
    for i in range(1, 6):
        style(f"{get_column_letter(i)}{row}", border=True)
    # Example formula: Amount = Qty * Rate
    style(f"E{row}", value=f'=IF(OR(C{row}="",D{row}=""),"",C{row}*D{row})',
          align="right", border=True, number_format=money)

# Total
style(f"D{LAST+1}", value="TOTAL", bold=True, align="right", color=WHITE, fill=BRAND, border=True)
style(f"E{LAST+1}", value=f"=SUM(E{FIRST}:E{LAST})", bold=True, align="right",
      color=WHITE, fill=BRAND, border=True, number_format=money)

out = Path("public/files/__SLUG__.xlsx")
out.parent.mkdir(parents=True, exist_ok=True)
wb.save(out)
print(f"Saved: {out.resolve()}")
'''

SEED_BLOCK = '''
  const cat_SLUGUNDER_ = await prisma.templateCategory.findUnique({
    where: { slug: "__CATEGORY__" }
  });
  if (cat_SLUGUNDER_) {
    const data_SLUGUNDER_ = {
      title: "__TITLE__",
      description: "TODO one-sentence summary with the keyword.",
      content: "TODO 3-4 sentences explaining what it is and how it helps.",
      fileUrl: "/files/__SLUG__.xlsx",
      thumbnailUrl: null,
      categoryId: cat_SLUGUNDER_.id,
      seoTitle: "TODO keyword-rich title | ExcelHub",
      seoDescription: "TODO 1-2 sentences with the keyword, under 160 chars.",
      features: ["TODO", "TODO", "TODO", "TODO"],
      instructions: ["TODO", "TODO", "TODO", "TODO"],
      faqs: [
        { question: "TODO?", answer: "TODO." },
        { question: "TODO?", answer: "TODO." }
      ]
    };
    await prisma.template.upsert({
      where: { slug: "__SLUG__" },
      update: { fileUrl: data_SLUGUNDER_.fileUrl },
      create: { slug: "__SLUG__", ...data_SLUGUNDER_ }
    });
  }
'''


def main():
    print("\n=== New ExcelHub template scaffolder ===\n")
    keyword = input("Target keyword (e.g. 'salary slip template excel'): ").strip()
    title = input("Template title (e.g. 'Salary Slip Template'): ").strip()
    if not title:
        print("A title is required. Aborting.")
        return
    slug = slugify(title)
    slug_under = slug.replace("-", "_")

    print("\nCategories:")
    for c in CATEGORIES:
        print(f"  - {c}")
    category = input("\nCategory slug [accounting-templates]: ").strip() or "accounting-templates"
    if category not in CATEGORIES:
        print(f"  '{category}' is not a known category; using business-templates.")
        category = "business-templates"

    def fill(s):
        return (s.replace("__TITLE__", title)
                 .replace("__SLUGUNDER__", slug_under)
                 .replace("__SLUG__", slug)
                 .replace("__CATEGORY__", category)
                 .replace("_SLUGUNDER_", slug_under))

    gen_path = Path(f"scripts/generate_{slug_under}.py")
    if gen_path.exists():
        print(f"\n{gen_path} already exists — not overwriting.")
    else:
        gen_path.write_text(fill(GENERATOR_SKELETON))
        print(f"\n✓ Created {gen_path}")

    print("\n--- Keyword (for your notes) ---")
    print(f"  {keyword or '(none entered)'}")

    print("\n========================================================")
    print("NEXT STEPS")
    print("========================================================")
    print(f"1. Edit {gen_path} — design your real columns + formulas.")
    print(f"2. Run it:  .venv-tools/bin/python {gen_path}")
    print(f"   -> creates public/files/{slug}.xlsx (open it and test the formulas)")
    print("3. Paste the block below into prisma/seed.ts, just ABOVE the line")
    print("   'if (!seedDemo) {', and fill in every TODO (use the AI prompt in")
    print("   PLAYBOOK.md to write the content).")
    print("4. Deploy:  vercel --prod")
    print(f"5. In Search Console, request indexing for:")
    print(f"   https://excelhub-omega.vercel.app/templates/{slug}")
    print("\n--------- paste this into prisma/seed.ts ---------")
    print(fill(SEED_BLOCK))
    print("--------------------------------------------------")


if __name__ == "__main__":
    main()

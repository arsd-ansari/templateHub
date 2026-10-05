"""
Generates an RCM invoice format (.xlsx) — the Rule 46 tax invoice a
registered supplier issues when the recipient pays the GST.

Sample: a firm of advocates bills a business entity for legal services.
Taxable fee ₹40,000 at 18%. Intra-state: CGST ₹3,600 + SGST ₹3,600.
That GST is not added to the fee the firm collects.

This is not a self invoice and not the GTA consignment bill.

Run:  .venv-tools/bin/python scripts/generate_rcm_invoice.py
Output: public/files/rcm-tax-invoice-format.xlsx
"""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

BRAND = "0F766E"
BRAND_LIGHT = "D7EEEB"
GREY = "F1F5F9"
WHITE = "FFFFFF"
DARK = "16181D"
ACCENT = "B45309"

thin = Side(style="thin", color="C9D2DC")
box = Border(left=thin, right=thin, top=thin, bottom=thin)

money_fmt = "#,##0.00"
pct_fmt = "0%"

wb = Workbook()
ws = wb.active
ws.title = "RCM Invoice"
ws.sheet_view.showGridLines = False
ws.page_setup.orientation = "landscape"
ws.page_setup.fitToPage = True
ws.page_setup.fitToWidth = 1
ws.page_setup.fitToHeight = 1
ws.page_setup.paperSize = ws.PAPERSIZE_A4
ws.sheet_properties.pageSetUpPr.fitToPage = True
ws.print_title_rows = "1:2"
ws.page_setup.horizontalCentered = True
ws.oddHeader.left.text = "RCM Invoice — tax payable by recipient"
ws.oddFooter.right.text = "Page &P of &N"
ws.page_margins.left = 0.4
ws.page_margins.right = 0.4
ws.page_margins.top = 0.6
ws.page_margins.bottom = 0.5

widths = [6, 42, 12, 8, 14, 16, 10, 14, 14, 14, 18]
for i, w in enumerate(widths, start=1):
    ws.column_dimensions[get_column_letter(i)].width = w


def style_cell(ref, *, value=None, bold=False, size=11, color=DARK,
               fill=None, align="left", border=False, wrap=False,
               number_format=None, italic=False):
    c = ws[ref]
    if value is not None:
        c.value = value
    c.font = Font(name="Calibri", bold=bold, size=size, color=color, italic=italic)
    c.alignment = Alignment(horizontal=align, vertical="center", wrap_text=wrap)
    if fill:
        c.fill = PatternFill("solid", fgColor=fill)
    if border:
        c.border = box
    if number_format:
        c.number_format = number_format
    return c


ws.merge_cells("A1:K1")
style_cell("A1", value="RCM TAX INVOICE  ·  TAX PAYABLE BY RECIPIENT", bold=True, size=20,
           color=WHITE, fill=BRAND, align="center")
ws.row_dimensions[1].height = 34

ws.merge_cells("A2:K2")
style_cell(
    "A2",
    value=("Issued by a registered supplier. Rule 46 requires this invoice to say whether tax is payable on reverse charge. "
           "Do not collect GST — the recipient pays it to the government. Set Supply type to Intra or Inter."),
    italic=True, size=10, color=ACCENT, align="center", wrap=True
)
ws.row_dimensions[2].height = 32

ws.merge_cells("A4:E4")
style_cell("A4", value="SUPPLIER (Registered — issuer of this invoice)",
           bold=True, color=WHITE, fill=BRAND, border=True)
for row, key in ((5, "Legal Name"), (6, "Address"), (7, "GSTIN"), (8, "State & Code")):
    ws.merge_cells(f"A{row}:E{row}")
    style_cell(f"A{row}", value=f"{key}: [ ]", border=True, wrap=True)

# H6 is the supply-type switch. Formulas read it. Do not merge H6.
meta = [
    (4, "Invoice No.", "[RCM-INV-001]"),
    (5, "Invoice Date", "[DD-MM-YYYY]"),
    (6, "Supply type", "Intra"),
    (7, "Place of Supply", "[Recipient's State (Code)]"),
    (8, "Reverse Charge", "Yes"),
]
for row, label, val in meta:
    style_cell(f"G{row}", value=label, bold=True, fill=GREY, border=True)
    if row == 6:
        style_cell(f"H{row}", value=val, border=True, bold=True, color=ACCENT, align="center")
        ws.merge_cells(f"I{row}:K{row}")
        style_cell(f"I{row}", value="Type Intra or Inter", border=True, italic=True, size=10, color=ACCENT)
    else:
        ws.merge_cells(f"H{row}:K{row}")
        style_cell(
            f"H{row}", value=val, border=True,
            bold=(label == "Reverse Charge"),
            color=(ACCENT if label == "Reverse Charge" else DARK)
        )

supply_validation = DataValidation(type="list", formula1='"Intra,Inter"', allow_blank=False)
supply_validation.error = "Enter Intra or Inter"
supply_validation.errorTitle = "Supply type"
supply_validation.prompt = "Intra splits GST into CGST and SGST. Inter charges IGST."
supply_validation.promptTitle = "Supply type"
supply_validation.add("H6")
ws.add_data_validation(supply_validation)

ws.merge_cells("A10:E10")
style_cell("A10", value="RECIPIENT (business entity — pays GST to the government)",
           bold=True, color=WHITE, fill=BRAND, border=True)
for row, key in ((11, "Legal Name"), (12, "Address"), (13, "GSTIN"), (14, "State & Code")):
    ws.merge_cells(f"A{row}:E{row}")
    style_cell(f"A{row}", value=f"{key}: [ ]", border=True, wrap=True)

ws.merge_cells("F10:K10")
style_cell("F10", value="SAMPLE — FIRM OF ADVOCATES, LEGAL SERVICES", bold=True, color=WHITE, fill=BRAND, border=True)
ws.merge_cells("F11:K14")
style_cell(
    "F11",
    value=("Intra-state sample: taxable fee ₹40,000 at 18% = CGST ₹3,600 + SGST ₹3,600. "
           "IGST is 0. The firm collects ₹40,000. The recipient pays ₹7,200 to the government. "
           "Switch H6 to Inter and the same fee becomes IGST ₹7,200, with CGST and SGST at 0."),
    border=True, wrap=True, italic=True, size=10, color=ACCENT
)
ws.row_dimensions[11].height = 22
ws.row_dimensions[12].height = 20
ws.row_dimensions[13].height = 20
ws.row_dimensions[14].height = 20

HEADER_ROW = 16
headers = [
    "S.No", "Description of Goods/Services", "SAC", "Qty", "Rate",
    "Taxable Value", "GST %", "CGST", "SGST", "IGST", "Payable to supplier"
]
for i, h in enumerate(headers, start=1):
    col = get_column_letter(i)
    style_cell(f"{col}{HEADER_ROW}", value=h, bold=True, color=WHITE,
               fill=BRAND, align="center", border=True, wrap=True)
ws.row_dimensions[HEADER_ROW].height = 30

FIRST = HEADER_ROW + 1
N_ROWS = 6
LAST = FIRST + N_ROWS - 1

samples = [
    (1, "Legal services — drafting and advisory on a commercial agreement (firm of advocates)", "998212", 1, 40000, 0.18),
]

for idx, row in enumerate(range(FIRST, LAST + 1)):
    sample = samples[idx] if idx < len(samples) else None
    style_cell(f"A{row}", value=(sample[0] if sample else None), align="center", border=True)
    style_cell(f"B{row}", value=(sample[1] if sample else None), border=True, wrap=True)
    style_cell(f"C{row}", value=(sample[2] if sample else None), align="center", border=True)
    style_cell(f"D{row}", value=(sample[3] if sample else None), align="center", border=True)
    style_cell(f"E{row}", value=(sample[4] if sample else None), align="right", border=True, number_format=money_fmt)
    style_cell(f"G{row}", value=(sample[5] if sample else None), align="center", border=True, number_format=pct_fmt)
    style_cell(
        f"F{row}",
        value=f'=IF(OR(D{row}="",E{row}=""),"",ROUND(D{row}*E{row},2))',
        align="right", border=True, number_format=money_fmt
    )
    # Intra (anything other than Inter): split the rate. Inter: IGST at the full rate.
    style_cell(
        f"H{row}",
        value=f'=IF(F{row}="","",IF(UPPER(TRIM($H$6))="INTER",0,ROUND(F{row}*G{row}/2,2)))',
        align="right", border=True, number_format=money_fmt
    )
    style_cell(
        f"I{row}",
        value=f'=IF(F{row}="","",IF(UPPER(TRIM($H$6))="INTER",0,ROUND(F{row}*G{row}/2,2)))',
        align="right", border=True, number_format=money_fmt
    )
    style_cell(
        f"J{row}",
        value=f'=IF(F{row}="","",IF(UPPER(TRIM($H$6))="INTER",ROUND(F{row}*G{row},2),0))',
        align="right", border=True, number_format=money_fmt
    )
    style_cell(
        f"K{row}",
        value=f'=IF(F{row}="","",F{row})',
        align="right", border=True, number_format=money_fmt
    )
    if idx % 2 == 1:
        for i in range(1, 12):
            ws[f"{get_column_letter(i)}{row}"].fill = PatternFill("solid", fgColor=GREY)
    ws.row_dimensions[row].height = 20
ws.row_dimensions[FIRST].height = 32


def total_row(row, label, formula, *, strong=False, accent=False):
    ws.merge_cells(f"A{row}:J{row}")
    fill = BRAND if strong else ("FEF3C7" if accent else BRAND_LIGHT)
    color = WHITE if strong else (ACCENT if accent else DARK)
    style_cell(f"A{row}", value=label, bold=True, align="right", fill=fill, color=color, border=True)
    style_cell(
        f"K{row}", value=formula, bold=True, align="right", border=True, number_format=money_fmt,
        fill=fill, color=color, size=(12 if strong else 11)
    )


base = LAST + 1
total_row(base, "Amount payable to supplier (do not add GST)", f"=ROUND(SUM(K{FIRST}:K{LAST}),2)", strong=True)
total_row(base + 1, "CGST payable by recipient to government", f"=ROUND(SUM(H{FIRST}:H{LAST}),2)", accent=True)
total_row(base + 2, "SGST payable by recipient to government", f"=ROUND(SUM(I{FIRST}:I{LAST}),2)", accent=True)
total_row(base + 3, "IGST payable by recipient to government", f"=ROUND(SUM(J{FIRST}:J{LAST}),2)", accent=True)
total_row(base + 4, "Total GST under reverse charge (not collected on this invoice)", f"=ROUND(K{base+1}+K{base+2}+K{base+3},2)", accent=True)

words = base + 6
ws.merge_cells(f"A{words}:K{words}")
style_cell(
    f"A{words}",
    value="Amount payable to supplier, in words: [ Rupees ______________________ Only ]",
    bold=True, italic=True
)

decl = words + 2
ws.merge_cells(f"A{decl}:K{decl}")
style_cell(f"A{decl}", value="DECLARATION", bold=True, color=WHITE, fill=BRAND, border=True, align="center")
lines = [
    "Tax is payable on reverse charge basis by the recipient. The supplier has not collected GST on this invoice.",
    "Supply type Intra: CGST and SGST are each half the GST %. Supply type Inter: IGST is the full GST %, and CGST and SGST are 0.",
    "Sample line is a firm of advocates billing legal services to a business entity (Notification 13/2017-Central Tax (Rate)). It is not a self invoice.",
    "If the supplier is not registered, do not use this invoice. The recipient issues a self invoice under Section 31(3)(f) instead.",
]
for i, line in enumerate(lines):
    ws.merge_cells(f"A{decl+1+i}:K{decl+1+i}")
    style_cell(f"A{decl+1+i}", value=line, border=True, size=10, wrap=True)
    ws.row_dimensions[decl + 1 + i].height = 28

sign = decl + len(lines) + 2
ws.merge_cells(f"H{sign}:K{sign}")
style_cell(f"H{sign}", value="For [Supplier Name]", bold=True, align="right")
ws.merge_cells(f"H{sign+2}:K{sign+2}")
style_cell(f"H{sign+2}", value="Authorised Signatory", align="right", color=BRAND)

ws.freeze_panes = f"A{FIRST}"
ws.print_area = f"A1:K{sign+2}"

out = Path("public/files/rcm-tax-invoice-format.xlsx")
out.parent.mkdir(parents=True, exist_ok=True)
wb.save(out)
print(f"Saved: {out.resolve()}")

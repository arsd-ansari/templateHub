"""
Generates a URD purchase bill (.xlsx).

An unregistered supplier cannot issue a tax invoice (Section 31 applies to a
registered person) and must not collect GST (Section 32). This bill is the
commercial record of that purchase. It is not a self invoice.

Sample lines (stationery bought from an unregistered shop):
  20 × ₹180  copier paper, ream     HSN 4802 = ₹3,600
  50 × ₹10   ball pens              HSN 9608 = ₹500
   5 × ₹140  desktop stapler        HSN 8472 = ₹700
  Amount payable = ₹4,800. GST is not charged.

Run:  .venv-tools/bin/python scripts/generate_urd_purchase_bill.py
Output: public/files/urd-purchase-bill-format.xlsx
"""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

BRAND = "0F766E"
BRAND_LIGHT = "D7EEEB"
GREY = "F1F5F9"
WHITE = "FFFFFF"
DARK = "16181D"
ACCENT = "B45309"
AMBER = "FEF3C7"

thin = Side(style="thin", color="C9D2DC")
box = Border(left=thin, right=thin, top=thin, bottom=thin)
money_fmt = "#,##0.00"

wb = Workbook()
ws = wb.active
ws.title = "URD Purchase Bill"
ws.sheet_view.showGridLines = False
ws.page_setup.orientation = "landscape"
ws.page_setup.fitToPage = True
ws.page_setup.fitToWidth = 1
ws.page_setup.fitToHeight = 1
ws.page_setup.paperSize = ws.PAPERSIZE_A4
ws.sheet_properties.pageSetUpPr.fitToPage = True
ws.print_title_rows = "1:2"
ws.page_setup.horizontalCentered = True
ws.oddHeader.left.text = "URD purchase bill — supplier is not registered — no GST"
ws.oddFooter.right.text = "Page &P of &N"
ws.page_margins.left = 0.4
ws.page_margins.right = 0.4
ws.page_margins.top = 0.6
ws.page_margins.bottom = 0.5

widths = [6, 46, 12, 10, 12, 14, 16]
for i, w in enumerate(widths, start=1):
    ws.column_dimensions[get_column_letter(i)].width = w


def style_cell(ref, *, value=None, bold=False, size=11, color=DARK,
               fill=None, align="left", border=False, wrap=False,
               number_format=None, italic=False):
    c = ws[ref]
    if value is not None:
        c.value = value
    c.font = Font(name="Calibri", bold=bold, size=size, color=color, italic=italic)
    c.alignment = Alignment(horizontal=align, vertical="center", wrap_text=wrap)
    if fill:
        c.fill = PatternFill("solid", fgColor=fill)
    if border:
        c.border = box
    if number_format:
        c.number_format = number_format
    return c


ws.merge_cells("A1:G1")
style_cell(
    "A1",
    value="URD PURCHASE BILL  ·  SUPPLIER NOT REGISTERED  ·  NO GST",
    bold=True, size=18, color=WHITE, fill=BRAND, align="center",
)
ws.row_dimensions[1].height = 34

ws.merge_cells("A2:G2")
style_cell(
    "A2",
    value=("Commercial bill from an unregistered supplier. Section 32: do not collect GST. "
           "This is not a tax invoice and not a self invoice. Sample total ₹4,800."),
    italic=True, size=10, color=ACCENT, align="center", wrap=True,
)
ws.row_dimensions[2].height = 32

ws.merge_cells("A4:D4")
style_cell("A4", value="SUPPLIER (unregistered — issuing this bill)",
           bold=True, color=WHITE, fill=BRAND, border=True)
for row, text in (
    (5, "Legal Name: [Shop or person]"),
    (6, "Address: [ ]"),
    (7, "GSTIN: Not registered — do not type a GSTIN"),
    (8, "State: [State name]"),
):
    ws.merge_cells(f"A{row}:D{row}")
    style_cell(f"A{row}", value=text, border=True, wrap=True,
               color=(ACCENT if row == 7 else DARK),
               bold=(row == 7))

meta = [
    (4, "Bill No.", "[UPB-001]"),
    (5, "Bill Date", "[DD-MM-YYYY]"),
    (6, "Place of supply", "[Buyer's state]"),
    (7, "Supplier status", "Unregistered"),
    (8, "GST on this bill", "Nil"),
]
for row, label, val in meta:
    style_cell(f"F{row}", value=label, bold=True, fill=GREY, border=True)
    style_cell(
        f"G{row}", value=val, border=True,
        bold=(label in ("Supplier status", "GST on this bill")),
        color=(ACCENT if label in ("Supplier status", "GST on this bill") else DARK),
        align="center", wrap=True,
    )
ws.row_dimensions[7].height = 30

ws.merge_cells("A10:D10")
style_cell("A10", value="BUYER (keeps this bill as the purchase record)",
           bold=True, color=WHITE, fill=BRAND, border=True)
for row, text in (
    (11, "Legal Name: [Registered business, if the buyer has a GSTIN]"),
    (12, "Address: [ ]"),
    (13, "GSTIN: [Buyer's GSTIN, or Unregistered]"),
    (14, "State: [State name]"),
):
    ws.merge_cells(f"A{row}:D{row}")
    style_cell(f"A{row}", value=text, border=True, wrap=True)

ws.merge_cells("F10:G10")
style_cell("F10", value="WHAT THIS BILL IS", bold=True, color=WHITE, fill=BRAND, border=True)
ws.merge_cells("F11:G14")
style_cell(
    "F11",
    value=("Sample: paper ₹3,600 + pens ₹500 + stapler ₹700 = ₹4,800. "
           "Amount payable equals that total. There is no GST column to fill."),
    border=True, wrap=True, italic=True, size=10, color=ACCENT,
)
ws.row_dimensions[11].height = 20
ws.row_dimensions[12].height = 20
ws.row_dimensions[13].height = 20
ws.row_dimensions[14].height = 22

HEADER_ROW = 16
headers = ["S.No", "Description of goods", "HSN", "Qty", "Unit", "Rate", "Amount"]
for i, h in enumerate(headers, start=1):
    col = get_column_letter(i)
    style_cell(f"{col}{HEADER_ROW}", value=h, bold=True, color=WHITE,
               fill=BRAND, align="center", border=True, wrap=True)
ws.row_dimensions[HEADER_ROW].height = 22

FIRST = 17
N_ROWS = 6
LAST = FIRST + N_ROWS - 1  # 22

samples = [
    (1, "Copier paper, 75 gsm, ream", "4802", 20, "Nos", 180),
    (2, "Ball pens", "9608", 50, "Nos", 10),
    (3, "Desktop stapler", "8472", 5, "Nos", 140),
]

for idx, row in enumerate(range(FIRST, LAST + 1)):
    sample = samples[idx] if idx < len(samples) else None
    style_cell(f"A{row}", value=(sample[0] if sample else None), align="center", border=True)
    style_cell(f"B{row}", value=(sample[1] if sample else None), border=True, wrap=True)
    style_cell(f"C{row}", value=(sample[2] if sample else None), align="center", border=True)
    style_cell(f"D{row}", value=(sample[3] if sample else None), align="center", border=True)
    style_cell(f"E{row}", value=(sample[4] if sample else None), align="center", border=True)
    style_cell(f"F{row}", value=(sample[5] if sample else None), align="right", border=True, number_format=money_fmt)
    style_cell(
        f"G{row}",
        value=f'=IF(OR(D{row}="",F{row}=""),"",ROUND(D{row}*F{row},2))',
        align="right", border=True, number_format=money_fmt,
    )
    if idx % 2 == 1:
        for i in range(1, 8):
            ws[f"{get_column_letter(i)}{row}"].fill = PatternFill("solid", fgColor=GREY)
    ws.row_dimensions[row].height = 22

base = LAST + 1  # 23
ws.merge_cells(f"A{base}:F{base}")
style_cell(
    f"A{base}",
    value="Value of goods",
    bold=True, align="right", fill=BRAND_LIGHT, border=True,
)
style_cell(
    f"G{base}",
    value=f"=ROUND(SUM(G{FIRST}:G{LAST}),2)",
    bold=True, align="right", border=True, number_format=money_fmt, fill=BRAND_LIGHT,
)

gst = base + 1  # 24
ws.merge_cells(f"A{gst}:F{gst}")
style_cell(
    f"A{gst}",
    value="GST (not collected — supplier is not registered)",
    bold=True, align="right", fill=AMBER, color=ACCENT, border=True,
)
style_cell(
    f"G{gst}",
    value=0,
    bold=True, align="right", border=True, number_format=money_fmt, fill=AMBER, color=ACCENT,
)

payable = base + 2  # 25
ws.merge_cells(f"A{payable}:F{payable}")
style_cell(
    f"A{payable}",
    value="Amount payable to the supplier",
    bold=True, align="right", fill=BRAND, color=WHITE, border=True,
)
style_cell(
    f"G{payable}",
    value=f"=G{base}+G{gst}",
    bold=True, align="right", border=True, number_format=money_fmt,
    fill=BRAND, color=WHITE, size=12,
)

words = payable + 2  # 27
ws.merge_cells(f"A{words}:G{words}")
style_cell(
    f"A{words}",
    value="Amount payable, in words: [ Rupees Four Thousand Eight Hundred Only ]",
    bold=True, italic=True,
)

decl = words + 2  # 29
ws.merge_cells(f"A{decl}:G{decl}")
style_cell(f"A{decl}", value="DECLARATION", bold=True, color=WHITE, fill=BRAND, border=True, align="center")
lines = [
    "The supplier is not registered under GST. This bill does not charge GST. Section 32 of the CGST Act says a person who is not a registered person shall not collect tax under the Act. Do not add CGST, SGST, or IGST to the amount payable.",
    "This is not a tax invoice under Section 31, and it is not a bill of supply under Section 31(3)(c). Those documents are issued by a registered person. A composition dealer is registered and uses a bill of supply, not this file.",
    "The buyer cannot claim input tax credit on this bill. No tax was charged. If the supply is notified under Section 9(3), or under Section 9(4) for a promoter, the registered buyer still issues a self invoice for the tax. This purchase bill does not replace that self invoice.",
    "Ordinary purchases from an unregistered shop are not reverse charge. Since 1 February 2019, Section 9(4) applies only to a notified class of recipient. A stationery bill to a normal business is not that class.",
]
for i, line in enumerate(lines):
    ws.merge_cells(f"A{decl + 1 + i}:G{decl + 1 + i}")
    style_cell(f"A{decl + 1 + i}", value=line, border=True, size=10, wrap=True)
    ws.row_dimensions[decl + 1 + i].height = 32

sign = decl + len(lines) + 2
ws.merge_cells(f"E{sign}:G{sign}")
style_cell(f"E{sign}", value="For [Supplier Name]", bold=True, align="right")
ws.merge_cells(f"E{sign + 2}:G{sign + 2}")
style_cell(f"E{sign + 2}", value="Authorised Signatory", align="right", color=BRAND)
ws.merge_cells(f"A{sign}:C{sign + 1}")
style_cell(
    f"A{sign}",
    value="HSN on this sample is for the buyer's purchase register. An unregistered supplier is not under the Rule 46 HSN duty.",
    italic=True, size=10, color=ACCENT, wrap=True,
)
ws.row_dimensions[sign].height = 20
ws.row_dimensions[sign + 1].height = 20

ws.freeze_panes = "A17"
ws.print_area = f"A1:G{sign + 2}"

expected = round(20 * 180 + 50 * 10 + 5 * 140, 2)
if expected != 4800:
    raise SystemExit(f"Sample arithmetic drifted: {expected}")
for row, qty, rate in ((17, 20, 180), (18, 50, 10), (19, 5, 140)):
    formula = ws[f"G{row}"].value
    if formula != f'=IF(OR(D{row}="",F{row}=""),"",ROUND(D{row}*F{row},2))':
        raise SystemExit(f"Unexpected formula in G{row}: {formula}")
    if ws[f"D{row}"].value != qty or ws[f"F{row}"].value != rate:
        raise SystemExit(f"Sample inputs drifted on row {row}")
if ws["G23"].value != "=ROUND(SUM(G17:G22),2)":
    raise SystemExit(f"Total formula drifted: {ws['G23'].value}")
if ws["G24"].value != 0:
    raise SystemExit("GST cell must stay 0")
if ws["G25"].value != "=G23+G24":
    raise SystemExit(f"Payable formula drifted: {ws['G25'].value}")

out = Path("public/files/urd-purchase-bill-format.xlsx")
out.parent.mkdir(parents=True, exist_ok=True)
wb.save(out)
print(f"Saved: {out.resolve()}")
print(f"Sample check: 20×180 + 50×10 + 5×140 = {expected:.2f}; GST 0; payable {expected:.2f}")

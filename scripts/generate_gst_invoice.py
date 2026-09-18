"""
Generates a professional, ready-to-use Indian GST Tax Invoice template (.xlsx)
with LIVE formulas — taxable value, CGST, SGST and totals auto-calculate.

Run:  .venv-tools/bin/python scripts/generate_gst_invoice.py
Output: public/files/gst-invoice-template.xlsx (shipped with the site)

This is an intra-state (CGST + SGST) invoice, the most common case for local
sales. For inter-state sales you would use IGST (full GST rate in one column).
"""

from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ---- Brand / style constants -------------------------------------------------
BRAND = "0F766E"        # teal (matches the ExcelHub site)
BRAND_LIGHT = "D7EEEB"
GREY = "F1F5F9"
WHITE = "FFFFFF"
DARK = "16181D"

thin = Side(style="thin", color="C9D2DC")
medium = Side(style="medium", color=BRAND)
box = Border(left=thin, right=thin, top=thin, bottom=thin)

money_fmt = '#,##0.00'
pct_fmt = '0%'

wb = Workbook()
ws = wb.active
ws.title = "Tax Invoice"
ws.sheet_view.showGridLines = False

# 10 columns: A..J
widths = [6, 26, 12, 8, 12, 14, 8, 13, 13, 14]
for i, w in enumerate(widths, start=1):
    ws.column_dimensions[get_column_letter(i)].width = w


def style_cell(ref, *, value=None, bold=False, size=11, color=DARK,
               fill=None, align="left", border=False, wrap=False,
               number_format=None, italic=False):
    c = ws[ref]
    if value is not None:
        c.value = value
    c.font = Font(name="Calibri", bold=bold, size=size, color=color, italic=italic)
    c.alignment = Alignment(horizontal=align, vertical="center", wrap_text=wrap)
    if fill:
        c.fill = PatternFill("solid", fgColor=fill)
    if border:
        c.border = box
    if number_format:
        c.number_format = number_format
    return c


# ---- Title -------------------------------------------------------------------
ws.merge_cells("A1:J1")
style_cell("A1", value="TAX INVOICE", bold=True, size=22, color=WHITE,
           fill=BRAND, align="center")
ws.row_dimensions[1].height = 34

# ---- Seller block ------------------------------------------------------------
ws.merge_cells("A3:E3")
style_cell("A3", value="[Your Company Name]", bold=True, size=14, color=BRAND)
ws.merge_cells("A4:E4")
style_cell("A4", value="[Street Address, City, State - PIN]")
ws.merge_cells("A5:E5")
style_cell("A5", value="GSTIN: [22AAAAA0000A1Z5]   |   State: [State Name] (Code: [00])")
ws.merge_cells("A6:E6")
style_cell("A6", value="Email: [you@company.com]   |   Phone: [+91 00000 00000]")

# ---- Invoice meta (right) ----------------------------------------------------
meta = [
    ("Invoice No.", "[INV-001]"),
    ("Invoice Date", "[DD-MM-YYYY]"),
    ("Place of Supply", "[State Name (Code)]"),
    ("Reverse Charge", "No"),
]
r = 3
for label, val in meta:
    style_cell(f"G{r}", value=label, bold=True, fill=GREY, border=True)
    ws.merge_cells(f"H{r}:J{r}")
    style_cell(f"H{r}", value=val, border=True)
    r += 1

# ---- Bill To / Ship To -------------------------------------------------------
ws.merge_cells("A8:E8")
style_cell("A8", value="BILL TO", bold=True, color=WHITE, fill=BRAND, border=True)
ws.merge_cells("F8:J8")
style_cell("F8", value="SHIP TO", bold=True, color=WHITE, fill=BRAND, border=True)

for row, key in ((9, "Customer Name"), (10, "Address"), (11, "GSTIN"),
                 (12, "State & Code")):
    ws.merge_cells(f"A{row}:E{row}")
    style_cell(f"A{row}", value=f"{key}: [ ]", border=True, wrap=True)
    ws.merge_cells(f"F{row}:J{row}")
    style_cell(f"F{row}", value=f"{key}: [ ]", border=True, wrap=True)

# ---- Line-item table header --------------------------------------------------
HEADER_ROW = 14
headers = ["S.No", "Item Description", "HSN/SAC", "Qty", "Rate",
           "Taxable Value", "GST %", "CGST", "SGST", "Total"]
for i, h in enumerate(headers, start=1):
    col = get_column_letter(i)
    style_cell(f"{col}{HEADER_ROW}", value=h, bold=True, color=WHITE,
               fill=BRAND, align="center", border=True, wrap=True)
ws.row_dimensions[HEADER_ROW].height = 28

# ---- Line-item rows (with live formulas) -------------------------------------
FIRST = HEADER_ROW + 1
N_ROWS = 8
LAST = FIRST + N_ROWS - 1

samples = [
    (1, "Wireless Mouse", "8471", 10, 450, 0.18),
    (2, "USB-C Cable 1m", "8544", 25, 120, 0.18),
    (3, "Laptop Stand (Aluminium)", "7616", 5, 899, 0.18),
]

for idx, row in enumerate(range(FIRST, LAST + 1)):
    sample = samples[idx] if idx < len(samples) else None
    # S.No / Description / HSN / Qty / Rate / GST%
    style_cell(f"A{row}", value=(sample[0] if sample else None),
               align="center", border=True)
    style_cell(f"B{row}", value=(sample[1] if sample else None), border=True)
    style_cell(f"C{row}", value=(sample[2] if sample else None),
               align="center", border=True)
    style_cell(f"D{row}", value=(sample[3] if sample else None),
               align="center", border=True)
    style_cell(f"E{row}", value=(sample[4] if sample else None),
               align="right", border=True, number_format=money_fmt)
    style_cell(f"G{row}", value=(sample[5] if sample else None),
               align="center", border=True, number_format=pct_fmt)
    # Taxable = Qty * Rate
    style_cell(f"F{row}",
               value=f'=IF(OR($D{row}="",$E{row}=""),"",ROUND($D{row}*$E{row},2))',
               align="right", border=True, number_format=money_fmt)
    # CGST = Taxable * GST% / 2
    style_cell(f"H{row}",
               value=f'=IF($F{row}="","",ROUND($F{row}*$G{row}/2,2))',
               align="right", border=True, number_format=money_fmt)
    # SGST = Taxable * GST% / 2
    style_cell(f"I{row}",
               value=f'=IF($F{row}="","",ROUND($F{row}*$G{row}/2,2))',
               align="right", border=True, number_format=money_fmt)
    # Total = Taxable + CGST + SGST
    style_cell(f"J{row}",
               value=f'=IF($F{row}="","",$F{row}+$H{row}+$I{row})',
               align="right", border=True, number_format=money_fmt)
    if idx % 2 == 1:
        for i in range(1, 11):
            ws[f"{get_column_letter(i)}{row}"].fill = PatternFill("solid", fgColor=GREY)

# ---- Totals ------------------------------------------------------------------
def total_row(row, label, formula, *, strong=False):
    ws.merge_cells(f"A{row}:G{row}")
    style_cell(f"A{row}", value=label, bold=True, align="right",
               fill=(BRAND if strong else BRAND_LIGHT),
               color=(WHITE if strong else DARK), border=True)
    ws.merge_cells(f"H{row}:I{row}")
    style_cell(f"H{row}", border=True, fill=(BRAND if strong else None))
    style_cell(f"J{row}", value=formula, bold=True, align="right",
               border=True, number_format=money_fmt,
               fill=(BRAND if strong else BRAND_LIGHT),
               color=(WHITE if strong else DARK), size=(12 if strong else 11))

tcg = LAST + 1
total_row(tcg, "Total Taxable Value", f"=ROUND(SUM(F{FIRST}:F{LAST}),2)")
total_row(tcg + 1, "Total CGST", f"=ROUND(SUM(H{FIRST}:H{LAST}),2)")
total_row(tcg + 2, "Total SGST", f"=ROUND(SUM(I{FIRST}:I{LAST}),2)")
total_row(tcg + 3, "Round Off",
          f"=ROUND(SUM(J{FIRST}:J{LAST}),0)-SUM(J{FIRST}:J{LAST})")
total_row(tcg + 4, "GRAND TOTAL (₹)",
          f"=ROUND(SUM(J{FIRST}:J{LAST}),0)", strong=True)

# ---- Amount in words / notes -------------------------------------------------
words = tcg + 6
ws.merge_cells(f"A{words}:J{words}")
style_cell(f"A{words}", value="Amount in words: [ Rupees ______________________ Only ]",
           bold=True, italic=True)

bank = words + 2
style_cell(f"A{bank}", value="Bank Details", bold=True, color=BRAND)
for i, line in enumerate([
    "Account Name: [ ]   |   Account No: [ ]",
    "Bank & Branch: [ ]   |   IFSC: [ ]   |   UPI: [ ]",
]):
    ws.merge_cells(f"A{bank+1+i}:F{bank+1+i}")
    style_cell(f"A{bank+1+i}", value=line)

style_cell(f"G{bank}", value="Terms & Conditions", bold=True, color=BRAND)
for i, line in enumerate([
    "1. Goods once sold will not be taken back.",
    "2. Payment due within 15 days of invoice date.",
]):
    ws.merge_cells(f"G{bank+1+i}:J{bank+1+i}")
    style_cell(f"G{bank+1+i}", value=line, size=10)

sign = bank + 4
ws.merge_cells(f"G{sign}:J{sign}")
style_cell(f"G{sign}", value="For [Your Company Name]", bold=True, align="right")
ws.merge_cells(f"G{sign+2}:J{sign+2}")
style_cell(f"G{sign+2}", value="Authorised Signatory", align="right", color=BRAND)

# Freeze the header so the table scrolls cleanly
ws.freeze_panes = f"A{FIRST}"

# ---- Save --------------------------------------------------------------------
out = Path("public/files/gst-invoice-template.xlsx")
out.parent.mkdir(parents=True, exist_ok=True)
wb.save(out)
print(f"Saved: {out.resolve()}")

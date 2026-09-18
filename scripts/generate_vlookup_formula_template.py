"""
Generates a VLOOKUP Formula Practice Template (.xlsx).

The workbook includes:
  - A product price lookup table with working exact-match VLOOKUP formulas.
  - An employee lookup table with department, location, and manager formulas.
  - A quick reference sheet explaining syntax and common mistakes.

Run:  .venv-tools/bin/python scripts/generate_vlookup_formula_template.py
Output: public/files/vlookup-formula-template.xlsx
"""

from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

BRAND = "0F766E"
BRAND_LIGHT = "D7EEEB"
GREY = "F1F5F9"
WHITE = "FFFFFF"
DARK = "16181D"
MUTED = "667085"
MONEY = '#,##0.00'

thin = Side(style="thin", color="C9D2DC")
box = Border(left=thin, right=thin, top=thin, bottom=thin)


def style(ws, ref, *, value=None, bold=False, size=11, color=DARK, fill=None,
          align="left", border=False, wrap=False, number_format=None, italic=False):
    c = ws[ref]
    if value is not None:
        c.value = value
    c.font = Font(name="Calibri", bold=bold, size=size, color=color, italic=italic)
    c.alignment = Alignment(horizontal=align, vertical="center", wrap_text=wrap)
    if fill:
        c.fill = PatternFill("solid", fgColor=fill)
    if border:
        c.border = box
    if number_format:
        c.number_format = number_format
    return c


def header_row(ws, row, headers):
    for i, h in enumerate(headers, start=1):
        style(ws, f"{get_column_letter(i)}{row}", value=h, bold=True, color=WHITE,
              fill=BRAND, align="center", border=True)
    ws.row_dimensions[row].height = 24


wb = Workbook()

# ============================ Sheet 1: Product Lookup ========================
pl = wb.active
pl.title = "Product Lookup"
pl.sheet_view.showGridLines = False
for i, width in enumerate([15, 24, 18, 14, 14, 16, 16], start=1):
    pl.column_dimensions[get_column_letter(i)].width = width

pl.merge_cells("A1:G1")
style(pl, "A1", value="VLOOKUP PRODUCT PRICE EXAMPLE", bold=True, size=20,
      color=WHITE, fill=BRAND, align="center")
pl.row_dimensions[1].height = 32
pl.merge_cells("A2:G2")
style(pl, "A2", value="Enter a product code in column A. Product name, category, price, and stock are returned with exact-match VLOOKUP formulas.",
      color=MUTED, wrap=True)

header_row(pl, 4, ["Code", "Product Name", "Category", "Price", "Stock", "Quantity", "Line Total"])
products = [
    ("P-1001", "Wireless Mouse", "Accessories", 699, 42),
    ("P-1002", "USB-C Cable", "Accessories", 299, 120),
    ("P-1003", "Laptop Stand", "Office", 1499, 18),
    ("P-1004", "Notebook", "Stationery", 99, 250),
    ("P-1005", "Desk Lamp", "Office", 899, 33),
]

lookup_start = 18
lookup_end = lookup_start + len(products) - 1

sample_codes = ["P-1001", "P-1003", "P-1005", "", "", ""]
for idx, row in enumerate(range(5, 11)):
    code = sample_codes[idx]
    style(pl, f"A{row}", value=code, border=True, align="center")
    style(pl, f"B{row}", value=f'=IFERROR(VLOOKUP($A{row},$A${lookup_start}:$E${lookup_end},2,FALSE),"")',
          border=True)
    style(pl, f"C{row}", value=f'=IFERROR(VLOOKUP($A{row},$A${lookup_start}:$E${lookup_end},3,FALSE),"")',
          border=True)
    style(pl, f"D{row}", value=f'=IFERROR(VLOOKUP($A{row},$A${lookup_start}:$E${lookup_end},4,FALSE),"")',
          border=True, align="right", number_format=MONEY)
    style(pl, f"E{row}", value=f'=IFERROR(VLOOKUP($A{row},$A${lookup_start}:$E${lookup_end},5,FALSE),"")',
          border=True, align="center")
    style(pl, f"F{row}", value=(2 if code else None), border=True, align="center")
    style(pl, f"G{row}", value=f'=IF(OR($D{row}="",$F{row}=""),"",ROUND($D{row}*$F{row},2))',
          border=True, align="right", number_format=MONEY)
    if idx % 2 == 1:
        for col in range(1, 8):
            pl[f"{get_column_letter(col)}{row}"].fill = PatternFill("solid", fgColor=GREY)

dv = DataValidation(type="list", formula1='"%s"' % ",".join([p[0] for p in products]), allow_blank=True)
pl.add_data_validation(dv)
dv.add("A5:A10")

pl.merge_cells("A13:G13")
style(pl, "A13", value="Formula used", bold=True, color=BRAND)
pl.merge_cells("A14:G14")
style(pl, "A14", value='=IFERROR(VLOOKUP($A5,$A$18:$E$22,2,FALSE),"")',
      italic=True, fill=BRAND_LIGHT, border=True)
pl.merge_cells("A15:G15")
style(pl, "A15", value="FALSE means exact match. The lookup code must be in the first column of the lookup table.",
      color=MUTED, wrap=True)

header_row(pl, 17, ["Code", "Product Name", "Category", "Price", "Stock"])
for idx, product in enumerate(products, start=lookup_start):
    for col, value in enumerate(product, start=1):
        style(pl, f"{get_column_letter(col)}{idx}", value=value, border=True,
              align=("right" if col == 4 else "center" if col in (1, 5) else "left"),
              number_format=(MONEY if col == 4 else None))

pl.freeze_panes = "A5"

# ============================ Sheet 2: Employee Lookup =======================
el = wb.create_sheet("Employee Lookup")
el.sheet_view.showGridLines = False
for i, width in enumerate([16, 24, 18, 18, 22], start=1):
    el.column_dimensions[get_column_letter(i)].width = width

el.merge_cells("A1:E1")
style(el, "A1", value="VLOOKUP EMPLOYEE DETAILS EXAMPLE", bold=True, size=18,
      color=WHITE, fill=BRAND, align="center")
el.merge_cells("A2:E2")
style(el, "A2", value="Type an employee ID and VLOOKUP returns the matching name, department, location, and manager.",
      color=MUTED)

header_row(el, 4, ["Employee ID", "Name", "Department", "Location", "Manager"])
employees = [
    ("E-001", "Asha Mehta", "Finance", "Mumbai", "Rohan Shah"),
    ("E-002", "Karan Iyer", "Sales", "Bengaluru", "Neha Rao"),
    ("E-003", "Nisha Kapoor", "HR", "Delhi", "Meera Singh"),
    ("E-004", "Vikram Nair", "Operations", "Pune", "Amit Jain"),
    ("E-005", "Sara Khan", "Support", "Hyderabad", "Neha Rao"),
]
emp_lookup_start = 16
emp_lookup_end = emp_lookup_start + len(employees) - 1

for idx, row in enumerate(range(5, 10)):
    emp_id = employees[idx][0] if idx < 3 else ""
    style(el, f"A{row}", value=emp_id, border=True, align="center")
    for col in range(2, 6):
        style(el, f"{get_column_letter(col)}{row}",
              value=f'=IFERROR(VLOOKUP($A{row},$A${emp_lookup_start}:$E${emp_lookup_end},{col},FALSE),"")',
              border=True)

emp_dv = DataValidation(type="list", formula1='"%s"' % ",".join([e[0] for e in employees]), allow_blank=True)
el.add_data_validation(emp_dv)
emp_dv.add("A5:A9")

header_row(el, 15, ["Employee ID", "Name", "Department", "Location", "Manager"])
for row, employee in enumerate(employees, start=emp_lookup_start):
    for col, value in enumerate(employee, start=1):
        style(el, f"{get_column_letter(col)}{row}", value=value, border=True,
              align=("center" if col == 1 else "left"))

# ============================ Sheet 3: Quick Guide ===========================
qg = wb.create_sheet("Quick Guide")
qg.sheet_view.showGridLines = False
for i, width in enumerate([26, 80], start=1):
    qg.column_dimensions[get_column_letter(i)].width = width

qg.merge_cells("A1:B1")
style(qg, "A1", value="VLOOKUP QUICK GUIDE", bold=True, size=20,
      color=WHITE, fill=BRAND, align="center")
qg.row_dimensions[1].height = 32

rows = [
    ("Syntax", "=VLOOKUP(lookup_value, table_array, col_index_num, [range_lookup])"),
    ("lookup_value", "The value you want to search for, such as a product code or employee ID."),
    ("table_array", "The lookup table. The lookup value must be in the first column of this selected range."),
    ("col_index_num", "The column number to return from the lookup table. Count from the left of table_array."),
    ("range_lookup", "Use FALSE for exact match in most business templates. TRUE is approximate match."),
    ("Best practice", "Use absolute references such as $A$18:$E$22 so the lookup range does not move when formulas are copied."),
    ("Common error", "#N/A means Excel did not find the lookup value in the first column of the lookup table."),
    ("Modern note", "In newer Excel versions, XLOOKUP is more flexible, but VLOOKUP is still common in shared workbooks."),
]

for row, (label, text) in enumerate(rows, start=3):
    style(qg, f"A{row}", value=label, bold=True, color=WHITE, fill=BRAND,
          border=True, wrap=True)
    style(qg, f"B{row}", value=text, border=True, wrap=True)
    qg.row_dimensions[row].height = 34

out = Path("public/files/vlookup-formula-template.xlsx")
out.parent.mkdir(parents=True, exist_ok=True)
wb.save(out)
print(f"Saved: {out.resolve()}")

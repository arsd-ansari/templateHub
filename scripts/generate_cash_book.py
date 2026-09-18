"""
Generates a Cash Book Template (.xlsx) — a chronological record of cash
received and paid, with an automatic running balance.

Run:  .venv-tools/bin/python scripts/generate_cash_book.py
Output: public/files/cash-book-template.xlsx (shipped with the site)
"""

from pathlib import Path
from datetime import date
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

BRAND = "0F766E"
BRAND_LIGHT = "D7EEEB"
GREY = "F1F5F9"
WHITE = "FFFFFF"
DARK = "16181D"
money = '#,##0.00'

thin = Side(style="thin", color="C9D2DC")
box = Border(left=thin, right=thin, top=thin, bottom=thin)

wb = Workbook()
ws = wb.active
ws.title = "Cash Book"
ws.sheet_view.showGridLines = False
for i, w in enumerate([14, 34, 14, 15, 15, 16], start=1):  # A..F
    ws.column_dimensions[get_column_letter(i)].width = w


def style(ref, *, value=None, bold=False, size=11, color=DARK, fill=None,
          align="left", border=False, number_format=None):
    c = ws[ref]
    if value is not None:
        c.value = value
    c.font = Font(name="Calibri", bold=bold, size=size, color=color)
    c.alignment = Alignment(horizontal=align, vertical="center")
    if fill:
        c.fill = PatternFill("solid", fgColor=fill)
    if border:
        c.border = box
    if number_format:
        c.number_format = number_format
    return c


# Title
ws.merge_cells("A1:F1")
style("A1", value="CASH BOOK", bold=True, size=22, color=WHITE, fill=BRAND, align="center")
ws.row_dimensions[1].height = 32
ws.merge_cells("A2:F2")
style("A2", value="Business: [Your Company Name]      Period: [Month / Year]", color="667085")

# Opening balance
ws.merge_cells("A3:E3")
style("A3", value="Opening Balance", bold=True, align="right", fill=BRAND_LIGHT, border=True)
style("F3", value=10000, bold=True, align="right", fill=BRAND_LIGHT, border=True, number_format=money)

# Header
H = 4
for i, h in enumerate(["Date", "Particulars", "Voucher No.", "Receipt (In)", "Payment (Out)", "Balance"], start=1):
    style(f"{get_column_letter(i)}{H}", value=h, bold=True, color=WHITE, fill=BRAND,
          align="center", border=True)
ws.row_dimensions[H].height = 24

FIRST = H + 1
N = 40
LAST = FIRST + N - 1

samples = [
    (date(2026, 6, 1), "Cash sales", "V-001", 5000, None),
    (date(2026, 6, 2), "Rent paid", "V-002", None, 8000),
    (date(2026, 6, 3), "Cash sales", "V-003", 3000, None),
]
for idx, row in enumerate(range(FIRST, LAST + 1)):
    s = samples[idx] if idx < len(samples) else None
    style(f"A{row}", value=(s[0] if s else None), border=True, align="center", number_format="dd-mm-yyyy")
    style(f"B{row}", value=(s[1] if s else None), border=True)
    style(f"C{row}", value=(s[2] if s else None), border=True, align="center")
    style(f"D{row}", value=(s[3] if s else None), border=True, align="right", number_format=money)
    style(f"E{row}", value=(s[4] if s else None), border=True, align="right", number_format=money)
    # Running balance = opening + cumulative receipts - cumulative payments (blank on empty rows)
    style(f"F{row}",
          value=f'=IF(AND(D{row}="",E{row}=""),"",$F$3+SUM($D${FIRST}:D{row})-SUM($E${FIRST}:E{row}))',
          border=True, align="right", number_format=money)
    if idx % 2 == 1:
        for i in range(1, 7):
            ws[f"{get_column_letter(i)}{row}"].fill = PatternFill("solid", fgColor=GREY)

# Totals
tot = LAST + 1
ws.merge_cells(f"A{tot}:C{tot}")
style(f"A{tot}", value="TOTAL", bold=True, align="right", color=WHITE, fill=BRAND, border=True)
style(f"D{tot}", value=f"=SUM(D{FIRST}:D{LAST})", bold=True, align="right", color=WHITE, fill=BRAND, border=True, number_format=money)
style(f"E{tot}", value=f"=SUM(E{FIRST}:E{LAST})", bold=True, align="right", color=WHITE, fill=BRAND, border=True, number_format=money)
style(f"F{tot}", color=WHITE, fill=BRAND, border=True)

clo = tot + 1
ws.merge_cells(f"A{clo}:E{clo}")
style(f"A{clo}", value="Closing Balance", bold=True, align="right", fill=BRAND_LIGHT, border=True)
style(f"F{clo}", value=f"=$F$3+SUM(D{FIRST}:D{LAST})-SUM(E{FIRST}:E{LAST})", bold=True,
      align="right", fill=BRAND_LIGHT, border=True, number_format=money, size=12)

ws.freeze_panes = f"A{FIRST}"

out = Path("public/files/cash-book-template.xlsx")
out.parent.mkdir(parents=True, exist_ok=True)
wb.save(out)
print(f"Saved: {out.resolve()}")

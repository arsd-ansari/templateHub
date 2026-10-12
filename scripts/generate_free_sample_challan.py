"""
Generates an invoice-for-free-sample delivery challan (.xlsx).

Free samples sent without consideration are not a supply (CBIC Circular
No. 92/11/2019-GST), except where Schedule I applies. Goods move on a
Rule 55 delivery challan. GST is not calculated and is not charged.

Sample lines (personal-care goods, not a drug sample):
  30 × ₹85  herbal shampoo 50 ml   HSN 3305 = ₹2,550
  40 × ₹60  moisturising cream 20 g HSN 3304 = ₹2,400
  24 × ₹35  bathing bar 75 g       HSN 3401 = ₹840
  Value of goods = ₹5,790, which is within the general ₹50,000 e-way bill limit.

Run:  .venv-tools/bin/python scripts/generate_free_sample_challan.py
Output: public/files/invoice-for-free-sample-under-gst.xlsx
"""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

BRAND = "0F766E"
BRAND_LIGHT = "D7EEEB"
GREY = "F1F5F9"
WHITE = "FFFFFF"
DARK = "16181D"
ACCENT = "B45309"
AMBER = "FEF3C7"

thin = Side(style="thin", color="C9D2DC")
box = Border(left=thin, right=thin, top=thin, bottom=thin)

money_fmt = '#,##0.00'

wb = Workbook()
ws = wb.active
ws.title = "Free Sample Challan"
ws.sheet_view.showGridLines = False
ws.page_setup.orientation = "landscape"
ws.page_setup.fitToPage = True
ws.page_setup.fitToWidth = 1
ws.page_setup.fitToHeight = 1
ws.page_setup.paperSize = ws.PAPERSIZE_A4
ws.sheet_properties.pageSetUpPr.fitToPage = True
ws.print_title_rows = "1:2"
ws.page_setup.horizontalCentered = True
ws.oddHeader.left.text = "Free sample delivery challan — not a tax invoice"
ws.oddFooter.right.text = "Page &P of &N"
ws.page_margins.left = 0.4
ws.page_margins.right = 0.4
ws.page_margins.top = 0.6
ws.page_margins.bottom = 0.5

widths = [6, 42, 12, 10, 12, 16, 18]
for i, w in enumerate(widths, start=1):
    ws.column_dimensions[get_column_letter(i)].width = w


def style_cell(ref, *, value=None, bold=False, size=11, color=DARK,
               fill=None, align="left", border=False, wrap=False,
               number_format=None, italic=False):
    c = ws[ref]
    if value is not None:
        c.value = value
    c.font = Font(name="Calibri", bold=bold, size=size, color=color, italic=italic)
    c.alignment = Alignment(horizontal=align, vertical="center", wrap_text=wrap)
    if fill:
        c.fill = PatternFill("solid", fgColor=fill)
    if border:
        c.border = box
    if number_format:
        c.number_format = number_format
    return c


ws.merge_cells("A1:G1")
style_cell(
    "A1",
    value="FREE SAMPLE DELIVERY CHALLAN  ·  NOT A TAX INVOICE",
    bold=True, size=18, color=WHITE, fill=BRAND, align="center"
)
ws.row_dimensions[1].height = 34

ws.merge_cells("A2:G2")
style_cell(
    "A2",
    value=("Rule 55 challan for goods sent as free samples without consideration. "
           "The consignee pays nothing. Do not add GST. Movement cell G6 is Intra or Inter."),
    italic=True, size=10, color=ACCENT, align="center", wrap=True
)
ws.row_dimensions[2].height = 32

ws.merge_cells("A4:D4")
style_cell("A4", value="CONSIGNER (registered — sending the samples)",
           bold=True, color=WHITE, fill=BRAND, border=True)
for row, key in ((5, "Legal Name"), (6, "Address"), (7, "GSTIN"), (8, "State & Code")):
    ws.merge_cells(f"A{row}:D{row}")
    style_cell(f"A{row}", value=f"{key}: [ ]", border=True, wrap=True)

# G6 is the movement switch. Do not merge G6.
meta = [
    (4, "Challan No.", "[DC-FS-001]"),
    (5, "Challan Date", "[DD-MM-YYYY]"),
    (6, "Movement", "Intra"),
    (7, "Place of supply", "[Only if Inter — state and code]"),
    (8, "Reason", "Free samples — not a supply"),
]
for row, label, val in meta:
    style_cell(f"F{row}", value=label, bold=True, fill=GREY, border=True)
    style_cell(
        f"G{row}", value=val, border=True,
        bold=(label == "Movement"),
        color=(ACCENT if label in ("Movement", "Reason") else DARK),
        align=("center" if label == "Movement" else "left"),
        wrap=True
    )
ws.row_dimensions[7].height = 30
ws.row_dimensions[8].height = 30

movement_validation = DataValidation(type="list", formula1='"Intra,Inter"', allow_blank=False)
movement_validation.error = "Enter Intra or Inter"
movement_validation.errorTitle = "Movement"
movement_validation.prompt = "Intra = same state. Inter = another state."
movement_validation.promptTitle = "Movement"
movement_validation.add("G6")
ws.add_data_validation(movement_validation)

ws.merge_cells("A10:D10")
style_cell("A10", value="CONSIGNEE (pays nothing for these goods)",
           bold=True, color=WHITE, fill=BRAND, border=True)
for row, key in (
    (11, "Name"),
    (12, "Address"),
    (13, "GSTIN (or Unregistered)"),
    (14, "State & Code"),
):
    ws.merge_cells(f"A{row}:D{row}")
    style_cell(f"A{row}", value=f"{key}: [ ]", border=True, wrap=True)

ws.merge_cells("F10:G10")
style_cell("F10", value="TRANSPORT", bold=True, color=WHITE, fill=BRAND, border=True)
for row, key in ((11, "Vehicle No."), (12, "Transporter")):
    style_cell(f"F{row}", value=key, bold=True, fill=GREY, border=True)
    style_cell(f"G{row}", value="[ ]", border=True)
ws.merge_cells("F13:G14")
style_cell(
    "F13",
    value=("Sample: shampoo, cream, and soap. "
           "Value = qty × value per unit. Total ₹5,790. No GST column."),
    border=True, wrap=True, italic=True, size=10, color=ACCENT
)
ws.row_dimensions[13].height = 22
ws.row_dimensions[14].height = 32

HEADER_ROW = 16
headers = ["S.No", "Description of goods", "HSN", "Qty", "Unit", "Value / unit", "Value"]
for i, h in enumerate(headers, start=1):
    col = get_column_letter(i)
    style_cell(f"{col}{HEADER_ROW}", value=h, bold=True, color=WHITE,
               fill=BRAND, align="center", border=True, wrap=True)
ws.row_dimensions[HEADER_ROW].height = 30

FIRST = HEADER_ROW + 1  # 17
N_ROWS = 6
LAST = FIRST + N_ROWS - 1  # 22

samples = [
    (1, "Herbal shampoo, 50 ml — free sample, not for sale", "3305", 30, "Nos", 85),
    (2, "Moisturising cream, 20 g — free sample, not for sale", "3304", 40, "Nos", 60),
    (3, "Bathing bar, 75 g — free sample, not for sale", "3401", 24, "Nos", 35),
]

for idx, row in enumerate(range(FIRST, LAST + 1)):
    sample = samples[idx] if idx < len(samples) else None
    style_cell(f"A{row}", value=(sample[0] if sample else None), align="center", border=True)
    style_cell(f"B{row}", value=(sample[1] if sample else None), border=True, wrap=True)
    style_cell(f"C{row}", value=(sample[2] if sample else None), align="center", border=True)
    style_cell(f"D{row}", value=(sample[3] if sample else None), align="center", border=True)
    style_cell(f"E{row}", value=(sample[4] if sample else None), align="center", border=True)
    style_cell(f"F{row}", value=(sample[5] if sample else None), align="right", border=True, number_format=money_fmt)
    style_cell(
        f"G{row}",
        value=f'=IF(OR(D{row}="",F{row}=""),"",ROUND(D{row}*F{row},2))',
        align="right", border=True, number_format=money_fmt
    )
    if idx % 2 == 1:
        for i in range(1, 8):
            ws[f"{get_column_letter(i)}{row}"].fill = PatternFill("solid", fgColor=GREY)
    ws.row_dimensions[row].height = 22
ws.row_dimensions[FIRST].height = 28
ws.row_dimensions[FIRST + 1].height = 28
ws.row_dimensions[FIRST + 2].height = 28

base = LAST + 1  # 23
ws.merge_cells(f"A{base}:F{base}")
style_cell(
    f"A{base}",
    value="Value of goods (the consignee does not pay this)",
    bold=True, align="right", fill=BRAND, color=WHITE, border=True
)
style_cell(
    f"G{base}",
    value=f"=ROUND(SUM(G{FIRST}:G{LAST}),2)",
    bold=True, align="right", border=True, number_format=money_fmt,
    fill=BRAND, color=WHITE, size=12
)

eway = base + 1  # 24
ws.merge_cells(f"A{eway}:F{eway}")
style_cell(
    f"A{eway}",
    value="E-way bill on the general ₹50,000 limit",
    bold=True, align="right", fill=AMBER, color=ACCENT, border=True
)
style_cell(
    f"G{eway}",
    value=f'=IF(G{base}="","",IF(G{base}>50000,"Required","Not required"))',
    bold=True, align="center", border=True, fill=AMBER, color=ACCENT
)

place = base + 2  # 25
ws.merge_cells(f"A{place}:F{place}")
style_cell(
    f"A{place}",
    value="Place of supply (Rule 55 — inter-state movement only)",
    bold=True, align="right", fill=BRAND_LIGHT, border=True
)
style_cell(
    f"G{place}",
    value=f'=IF(G6="Inter","Fill G7","Not required")',
    bold=True, align="center", border=True, fill=BRAND_LIGHT, color=DARK
)

words = base + 4  # 27
ws.merge_cells(f"A{words}:G{words}")
style_cell(
    f"A{words}",
    value="Value of goods, in words: [ Rupees Five Thousand Seven Hundred Ninety Only ]",
    bold=True, italic=True
)

decl = words + 2  # 29
ws.merge_cells(f"A{decl}:G{decl}")
style_cell(f"A{decl}", value="DECLARATION", bold=True, color=WHITE, fill=BRAND, border=True, align="center")
lines = [
    "These goods are free samples, sent without consideration. This is a delivery challan under Rule 55 of the CGST Rules for transportation for reasons other than supply. It is not a tax invoice under Section 31. No GST is charged.",
    "The value column is the value Rule 55 asks you to show. It is not a price the consignee pays. Input tax credit on goods disposed of as free samples is blocked by Section 17(5)(h) — reverse any credit already claimed (CBIC Circular No. 92/11/2019-GST).",
    "Do not use this challan where Schedule I treats the movement as a supply (for example a distinct person with another GSTIN, or a gift to an employee above ₹50,000 in a financial year). Issue a tax invoice and pay GST in that case.",
    "Challan number: 16 characters or fewer. Keep three copies — original for the consignee, duplicate for the transporter, triplicate for the consigner. Generate an e-way bill when the value exceeds the limit for that movement. This sheet flags the general ₹50,000 limit.",
]
for i, line in enumerate(lines):
    ws.merge_cells(f"A{decl+1+i}:G{decl+1+i}")
    style_cell(f"A{decl+1+i}", value=line, border=True, size=10, wrap=True)
    ws.row_dimensions[decl + 1 + i].height = 32

sign = decl + len(lines) + 2
ws.merge_cells(f"E{sign}:G{sign}")
style_cell(f"E{sign}", value="For [Consigner Name]", bold=True, align="right")
ws.merge_cells(f"E{sign+2}:G{sign+2}")
style_cell(f"E{sign+2}", value="Authorised Signatory", align="right", color=BRAND)
ws.merge_cells(f"A{sign}:C{sign}")
style_cell(f"A{sign}", value="Copy: Original for consignee", italic=True, size=10, color=ACCENT)
ws.merge_cells(f"A{sign+1}:C{sign+1}")
style_cell(f"A{sign+1}", value="(or Duplicate for transporter / Triplicate for consigner)", italic=True, size=10)

ws.freeze_panes = f"A{FIRST}"
ws.print_area = f"A1:G{sign+2}"

out = Path("public/files/invoice-for-free-sample-under-gst.xlsx")
out.parent.mkdir(parents=True, exist_ok=True)
wb.save(out)
print(f"Saved: {out.resolve()}")

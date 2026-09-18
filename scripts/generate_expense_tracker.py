"""
Generates a Business Expense Tracker template (.xlsx) with two linked sheets:
  - "Expenses": a log where you enter each expense (Date, Category dropdown,
    Description, Payment Method, Amount). Month auto-fills from the date.
  - "Summary": spend-by-category totals (live SUMIF), grand total, and a
    transaction count — all update automatically as you add rows.

Run:  .venv-tools/bin/python scripts/generate_expense_tracker.py
Output: public/files/business-expense-tracker.xlsx (shipped with the site)
"""

from pathlib import Path
from datetime import date
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

BRAND = "0F766E"
BRAND_LIGHT = "D7EEEB"
GREY = "F1F5F9"
WHITE = "FFFFFF"
DARK = "16181D"
money = '#,##0.00'

thin = Side(style="thin", color="C9D2DC")
box = Border(left=thin, right=thin, top=thin, bottom=thin)

CATEGORIES = ["Rent", "Utilities", "Salaries", "Inventory", "Marketing",
              "Travel", "Office Supplies", "Software", "Taxes", "Miscellaneous"]

FIRST = 4            # first data row on Expenses sheet
N_ROWS = 40
LAST = FIRST + N_ROWS - 1


def style(ws, ref, *, value=None, bold=False, size=11, color=DARK, fill=None,
          align="left", border=False, wrap=False, number_format=None):
    c = ws[ref]
    if value is not None:
        c.value = value
    c.font = Font(name="Calibri", bold=bold, size=size, color=color)
    c.alignment = Alignment(horizontal=align, vertical="center", wrap_text=wrap)
    if fill:
        c.fill = PatternFill("solid", fgColor=fill)
    if border:
        c.border = box
    if number_format:
        c.number_format = number_format
    return c


wb = Workbook()

# ============================ Sheet 1: Expenses ==============================
ex = wb.active
ex.title = "Expenses"
ex.sheet_view.showGridLines = False
for i, w in enumerate([14, 13, 18, 30, 16, 14], start=1):
    ex.column_dimensions[get_column_letter(i)].width = w

ex.merge_cells("A1:F1")
style(ex, "A1", value="BUSINESS EXPENSE TRACKER", bold=True, size=20,
      color=WHITE, fill=BRAND, align="center")
ex.row_dimensions[1].height = 32
ex.merge_cells("A2:F2")
style(ex, "A2", value="Enter each expense below. Month, totals and the Summary sheet update automatically.",
      color="667085")

headers = ["Date", "Month", "Category", "Description", "Payment Method", "Amount"]
for i, h in enumerate(headers, start=1):
    style(ex, f"{get_column_letter(i)}3", value=h, bold=True, color=WHITE,
          fill=BRAND, align="center", border=True)
ex.row_dimensions[3].height = 24

samples = [
    (date(2026, 6, 1), "Rent", "Office rent - June", "Bank Transfer", 25000),
    (date(2026, 6, 2), "Software", "Accounting software subscription", "Credit Card", 1499),
    (date(2026, 6, 3), "Travel", "Client meeting - cab", "Cash", 480),
]
for idx, row in enumerate(range(FIRST, LAST + 1)):
    s = samples[idx] if idx < len(samples) else None
    style(ex, f"A{row}", value=(s[0] if s else None), border=True, align="center",
          number_format="dd-mm-yyyy")
    # Month auto-derives from the Date
    style(ex, f"B{row}", value=f'=IF($A{row}="","",TEXT($A{row},"MMM YYYY"))',
          border=True, align="center")
    style(ex, f"C{row}", value=(s[1] if s else None), border=True)
    style(ex, f"D{row}", value=(s[2] if s else None), border=True)
    style(ex, f"E{row}", value=(s[3] if s else None), border=True, align="center")
    style(ex, f"F{row}", value=(s[4] if s else None), border=True, align="right",
          number_format=money)
    if idx % 2 == 1:
        for i in range(1, 7):
            ex[f"{get_column_letter(i)}{row}"].fill = PatternFill("solid", fgColor=GREY)

# Category dropdown on column C
dv = DataValidation(type="list", formula1='"%s"' % ",".join(CATEGORIES), allow_blank=True)
ex.add_data_validation(dv)
dv.add(f"C{FIRST}:C{LAST}")

# Total row
tot = LAST + 1
ex.merge_cells(f"A{tot}:E{tot}")
style(ex, f"A{tot}", value="TOTAL", bold=True, align="right", color=WHITE, fill=BRAND, border=True)
style(ex, f"F{tot}", value=f"=SUM(F{FIRST}:F{LAST})", bold=True, align="right",
      color=WHITE, fill=BRAND, border=True, number_format=money)

ex.freeze_panes = f"A{FIRST}"

# ============================ Sheet 2: Summary ===============================
sm = wb.create_sheet("Summary")
sm.sheet_view.showGridLines = False
for i, w in enumerate([22, 18], start=1):
    sm.column_dimensions[get_column_letter(i)].width = w

sm.merge_cells("A1:B1")
style(sm, "A1", value="EXPENSE SUMMARY", bold=True, size=18, color=WHITE,
      fill=BRAND, align="center")
sm.row_dimensions[1].height = 30

style(sm, "A3", value="Category", bold=True, color=WHITE, fill=BRAND, border=True)
style(sm, "B3", value="Total Spent", bold=True, color=WHITE, fill=BRAND, border=True, align="right")

r = 4
for cat in CATEGORIES:
    style(sm, f"A{r}", value=cat, border=True)
    style(sm, f"B{r}",
          value=f"=SUMIF(Expenses!$C${FIRST}:$C${LAST},$A{r},Expenses!$F${FIRST}:$F${LAST})",
          border=True, align="right", number_format=money)
    r += 1

style(sm, f"A{r}", value="GRAND TOTAL", bold=True, fill=BRAND_LIGHT, border=True)
style(sm, f"B{r}", value=f"=SUM(B4:B{r-1})", bold=True, fill=BRAND_LIGHT, border=True,
      align="right", number_format=money)
style(sm, f"A{r+1}", value="Transactions", bold=True, border=True)
style(sm, f"B{r+1}", value=f"=COUNT(Expenses!$F${FIRST}:$F${LAST})", bold=True,
      border=True, align="right")

out = Path("public/files/business-expense-tracker.xlsx")
out.parent.mkdir(parents=True, exist_ok=True)
wb.save(out)
print(f"Saved: {out.resolve()}")

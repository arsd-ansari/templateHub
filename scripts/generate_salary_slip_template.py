"""
Generates a professional Salary Slip Template (.xlsx) for Excel and Google Sheets.
Earnings, deductions, gross pay and net pay calculate automatically.

Run:  .venv-tools/bin/python scripts/generate_salary_slip_template.py
Output: public/files/salary-slip-template.xlsx (shipped with the site)
"""

from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

BRAND = "0F766E"
BRAND_LIGHT = "D7EEEB"
GREY = "F1F5F9"
WHITE = "FFFFFF"
DARK = "16181D"

money = '#,##0.00'
thin = Side(style="thin", color="C9D2DC")
box = Border(left=thin, right=thin, top=thin, bottom=thin)

wb = Workbook()
ws = wb.active
ws.title = "Salary Slip"
ws.sheet_view.showGridLines = False

for i, w in enumerate([18, 22, 18, 18, 18, 18], start=1):
    ws.column_dimensions[get_column_letter(i)].width = w


def style(ref, *, value=None, bold=False, size=11, color=DARK, fill=None,
          align="left", border=False, wrap=False, number_format=None, italic=False):
    cell = ws[ref]
    if value is not None:
        cell.value = value
    cell.font = Font(name="Calibri", bold=bold, size=size, color=color, italic=italic)
    cell.alignment = Alignment(horizontal=align, vertical="center", wrap_text=wrap)
    if fill:
        cell.fill = PatternFill("solid", fgColor=fill)
    if border:
        cell.border = box
    if number_format:
        cell.number_format = number_format
    return cell


ws.merge_cells("A1:F1")
style("A1", value="SALARY SLIP", bold=True, size=22, color=WHITE, fill=BRAND, align="center")
ws.row_dimensions[1].height = 34

ws.merge_cells("A3:F3")
style("A3", value="[Company Name]", bold=True, size=14, color=BRAND, align="center")
ws.merge_cells("A4:F4")
style("A4", value="[Company Address] | [Email] | [Phone]", align="center")

meta = [
    ("Salary Month", "[Month YYYY]", "Employee Name", "[Employee Name]"),
    ("Employee ID", "[EMP-001]", "Designation", "[Designation]"),
    ("Department", "[Department]", "Date of Joining", "[DD-MM-YYYY]"),
    ("Bank Account", "[Account No.]", "PAN / Tax ID", "[PAN / Tax ID]"),
]

row = 6
for left_label, left_value, right_label, right_value in meta:
    style(f"A{row}", value=left_label, bold=True, fill=GREY, border=True)
    style(f"B{row}", value=left_value, border=True)
    style(f"C{row}", value="", border=False)
    style(f"D{row}", value=right_label, bold=True, fill=GREY, border=True)
    ws.merge_cells(f"E{row}:F{row}")
    style(f"E{row}", value=right_value, border=True)
    row += 1

header_row = 12
ws.merge_cells(f"A{header_row}:C{header_row}")
style(f"A{header_row}", value="EARNINGS", bold=True, color=WHITE, fill=BRAND, align="center", border=True)
ws.merge_cells(f"D{header_row}:F{header_row}")
style(f"D{header_row}", value="DEDUCTIONS", bold=True, color=WHITE, fill=BRAND, align="center", border=True)

sub_header_row = header_row + 1
for ref, text in [
    ("A", "Component"), ("B", "Amount"), ("C", "Notes"),
    ("D", "Component"), ("E", "Amount"), ("F", "Notes"),
]:
    style(f"{ref}{sub_header_row}", value=text, bold=True, fill=BRAND_LIGHT, align="center", border=True)

earnings = [
    ("Basic Salary", 30000, "Monthly basic pay"),
    ("House Rent Allowance", 12000, "HRA"),
    ("Conveyance Allowance", 1600, "Transport"),
    ("Special Allowance", 6400, "Balancing allowance"),
    ("Bonus / Incentive", 0, "Optional"),
]
deductions = [
    ("Provident Fund", 3600, "Employee PF"),
    ("Professional Tax", 200, "If applicable"),
    ("Income Tax / TDS", 0, "If applicable"),
    ("Loan / Advance", 0, "Optional"),
    ("Other Deduction", 0, "Optional"),
]

first = sub_header_row + 1
for index in range(8):
    excel_row = first + index
    earning = earnings[index] if index < len(earnings) else ("", "", "")
    deduction = deductions[index] if index < len(deductions) else ("", "", "")
    style(f"A{excel_row}", value=earning[0], border=True)
    style(f"B{excel_row}", value=earning[1], align="right", border=True, number_format=money)
    style(f"C{excel_row}", value=earning[2], border=True)
    style(f"D{excel_row}", value=deduction[0], border=True)
    style(f"E{excel_row}", value=deduction[1], align="right", border=True, number_format=money)
    style(f"F{excel_row}", value=deduction[2], border=True)
    if index % 2 == 1:
        for col in "ABCDEF":
            ws[f"{col}{excel_row}"].fill = PatternFill("solid", fgColor=GREY)

last = first + 7
total_row = last + 1
style(f"A{total_row}", value="Gross Earnings", bold=True, color=WHITE, fill=BRAND, border=True)
style(f"B{total_row}", value=f"=SUM(B{first}:B{last})", bold=True, color=WHITE, fill=BRAND,
      align="right", border=True, number_format=money)
style(f"C{total_row}", fill=BRAND, border=True)
style(f"D{total_row}", value="Total Deductions", bold=True, color=WHITE, fill=BRAND, border=True)
style(f"E{total_row}", value=f"=SUM(E{first}:E{last})", bold=True, color=WHITE, fill=BRAND,
      align="right", border=True, number_format=money)
style(f"F{total_row}", fill=BRAND, border=True)

net_row = total_row + 2
ws.merge_cells(f"A{net_row}:D{net_row}")
style(f"A{net_row}", value="NET SALARY PAYABLE", bold=True, size=13, color=WHITE, fill=BRAND,
      align="right", border=True)
ws.merge_cells(f"E{net_row}:F{net_row}")
style(f"E{net_row}", value=f"=B{total_row}-E{total_row}", bold=True, size=13, color=WHITE,
      fill=BRAND, align="right", border=True, number_format=money)

words_row = net_row + 2
ws.merge_cells(f"A{words_row}:F{words_row}")
style(f"A{words_row}", value="Amount in words: [ Rupees ______________________ Only ]",
      bold=True, italic=True)

note_row = words_row + 2
ws.merge_cells(f"A{note_row}:F{note_row}")
style(f"A{note_row}", value="Note: This is a computer-generated salary slip. Please verify salary rules, tax deductions, and statutory contributions for your location before official use.",
      size=10, color="475569", wrap=True)

sign_row = note_row + 3
ws.merge_cells(f"A{sign_row}:C{sign_row}")
style(f"A{sign_row}", value="Employee Signature", color=BRAND)
ws.merge_cells(f"D{sign_row}:F{sign_row}")
style(f"D{sign_row}", value="Authorized Signatory", color=BRAND, align="right")

ws.freeze_panes = f"A{first}"

out = Path("public/files/salary-slip-template.xlsx")
out.parent.mkdir(parents=True, exist_ok=True)
wb.save(out)
print(f"Saved: {out.resolve()}")

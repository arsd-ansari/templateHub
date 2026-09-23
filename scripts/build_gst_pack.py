"""
Builds pack-only Excel extras and the paid GST zip.

Run:  /tmp/th-xlsx/bin/python scripts/build_gst_pack.py
      (or any python with openpyxl)

Output:
  content/packs/gst-invoice-igst-template.xlsx
  content/packs/gta-freight-worked-example.xlsx
  content/packs/rcm-document-register.xlsx
  content/packs/th-gst-pack-k8m2n4p6q8r0.zip
"""

from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

BRAND = "0F766E"
BRAND_LIGHT = "D7EEEB"
GREY = "F1F5F9"
WHITE = "FFFFFF"
DARK = "16181D"
ACCENT = "B45309"

thin = Side(style="thin", color="C9D2DC")
box = Border(left=thin, right=thin, top=thin, bottom=thin)
money_fmt = "#,##0.00"
pct_fmt = "0%"

ROOT = Path(__file__).resolve().parents[1]
PACK_DIR = ROOT / "content" / "packs"
ZIP_NAME = "th-gst-pack-k8m2n4p6q8r0.zip"


def style(ws, ref, *, value=None, bold=False, size=11, color=DARK,
          fill=None, align="left", border=False, wrap=False,
          number_format=None, italic=False):
    c = ws[ref]
    if value is not None:
        c.value = value
    c.font = Font(name="Calibri", bold=bold, size=size, color=color, italic=italic)
    c.alignment = Alignment(horizontal=align, vertical="center", wrap_text=wrap)
    if fill:
        c.fill = PatternFill("solid", fgColor=fill)
    if border:
        c.border = box
    if number_format:
        c.number_format = number_format
    return c


def write_igst_invoice(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Tax Invoice (IGST)"
    ws.sheet_view.showGridLines = False
    for i, w in enumerate([6, 28, 12, 8, 12, 14, 8, 14, 14], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ws.merge_cells("A1:I1")
    style(ws, "A1", value="TAX INVOICE (Inter-state — IGST)", bold=True, size=20,
          color=WHITE, fill=BRAND, align="center")
    ws.row_dimensions[1].height = 32
    ws.merge_cells("A2:I2")
    style(ws, "A2",
          value="Use this when place of supply is a different state from your GST registration. Do not split CGST+SGST.",
          italic=True, size=10, color=ACCENT, align="center")

    ws.merge_cells("A4:E4")
    style(ws, "A4", value="[Your Company Name]", bold=True, size=14, color=BRAND)
    ws.merge_cells("A5:E5")
    style(ws, "A5", value="GSTIN: [22AAAAA0000A1Z5]   |   State: [Your State] (Code: [00])")

    meta = [("Invoice No.", "[INV-001]"), ("Invoice Date", "[DD-MM-YYYY]"),
            ("Place of Supply", "[Other State (Code)]"), ("Reverse Charge", "No")]
    r = 4
    for label, val in meta:
        style(ws, f"G{r}", value=label, bold=True, fill=GREY, border=True)
        ws.merge_cells(f"H{r}:I{r}")
        style(ws, f"H{r}", value=val, border=True)
        r += 1

    ws.merge_cells("A9:I9")
    style(ws, "A9", value="BILL TO", bold=True, color=WHITE, fill=BRAND, border=True)
    for row, key in ((10, "Customer Name"), (11, "Address"), (12, "GSTIN"), (13, "State & Code")):
        ws.merge_cells(f"A{row}:I{row}")
        style(ws, f"A{row}", value=f"{key}: [ ]", border=True)

    header = 15
    for i, h in enumerate(["S.No", "Item Description", "HSN/SAC", "Qty", "Rate",
                           "Taxable Value", "GST %", "IGST", "Total"], start=1):
        style(ws, f"{get_column_letter(i)}{header}", value=h, bold=True, color=WHITE,
              fill=BRAND, align="center", border=True, wrap=True)
    first, n = header + 1, 8
    last = first + n - 1
    samples = [(1, "Consulting (inter-state)", "9983", 1, 25000, 0.18),
               (2, "Spare part — interstate", "8708", 4, 1800, 0.18)]
    for idx, row in enumerate(range(first, last + 1)):
        sample = samples[idx] if idx < len(samples) else None
        style(ws, f"A{row}", value=(sample[0] if sample else None), align="center", border=True)
        style(ws, f"B{row}", value=(sample[1] if sample else None), border=True)
        style(ws, f"C{row}", value=(sample[2] if sample else None), align="center", border=True)
        style(ws, f"D{row}", value=(sample[3] if sample else None), align="center", border=True)
        style(ws, f"E{row}", value=(sample[4] if sample else None), align="right",
              border=True, number_format=money_fmt)
        style(ws, f"G{row}", value=(sample[5] if sample else None), align="center",
              border=True, number_format=pct_fmt)
        style(ws, f"F{row}",
              value=f'=IF(OR($D{row}="",$E{row}=""),"",ROUND($D{row}*$E{row},2))',
              align="right", border=True, number_format=money_fmt)
        style(ws, f"H{row}",
              value=f'=IF($F{row}="","",ROUND($F{row}*$G{row},2))',
              align="right", border=True, number_format=money_fmt)
        style(ws, f"I{row}",
              value=f'=IF($F{row}="","",$F{row}+$H{row})',
              align="right", border=True, number_format=money_fmt)
        if idx % 2:
            for i in range(1, 10):
                ws[f"{get_column_letter(i)}{row}"].fill = PatternFill("solid", fgColor=GREY)

    def total(row, label, formula, strong=False):
        ws.merge_cells(f"A{row}:G{row}")
        style(ws, f"A{row}", value=label, bold=True, align="right",
              fill=(BRAND if strong else BRAND_LIGHT),
              color=(WHITE if strong else DARK), border=True)
        style(ws, f"H{row}", border=True, fill=(BRAND if strong else BRAND_LIGHT))
        style(ws, f"I{row}", value=formula, bold=True, align="right",
              border=True, number_format=money_fmt,
              fill=(BRAND if strong else BRAND_LIGHT),
              color=(WHITE if strong else DARK))

    t = last + 1
    total(t, "Total Taxable Value", f"=ROUND(SUM(F{first}:F{last}),2)")
    total(t + 1, "Total IGST", f"=ROUND(SUM(H{first}:H{last}),2)")
    total(t + 2, "GRAND TOTAL (₹)", f"=ROUND(SUM(I{first}:I{last}),0)", strong=True)
    ws.merge_cells(f"A{t+4}:I{t+4}")
    style(ws, f"A{t+4}",
          value="Amount in words: [ Rupees ______________________ Only ]",
          bold=True, italic=True)
    wb.save(path)


def write_gta_example(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "GTA RCM Example"
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 36
    ws.column_dimensions["B"].width = 28
    ws.column_dimensions["C"].width = 22

    ws.merge_cells("A1:C1")
    style(ws, "A1", value="WORKED EXAMPLE — GTA freight under RCM", bold=True,
          size=18, color=WHITE, fill=BRAND, align="center")
    ws.row_dimensions[1].height = 30
    ws.merge_cells("A2:C2")
    style(ws, "A2",
          value="Filled numbers for learning. Confirm the 5% vs 12% GTA rate with your CA.",
          italic=True, size=10, color=ACCENT, align="center")

    rows = [
        (4, "Self invoice no.", "SI-014"),
        (5, "Payment voucher no.", "PV-014"),
        (6, "Date of service (LR date)", "15-09-2026"),
        (7, "Recipient GSTIN (you)", "[Your GSTIN]"),
        (8, "GTA name", "Sample Roadways"),
        (9, "GTA GSTIN", "(blank — unregistered)"),
        (10, "Freight billed by GTA (₹)", 18500),
        (11, "GST rate", 0.05),
        (12, "CGST 2.5% (₹)", "=ROUND(B10*B11/2,2)"),
        (13, "SGST 2.5% (₹)", "=ROUND(B10*B11/2,2)"),
        (14, "GST to government (₹)", "=B12+B13"),
        (15, "Pay GTA (NEFT) (₹)", "=B10"),
        (16, "GSTR-3B Table 3.1(d) taxable", "=B10"),
        (17, "GSTR-3B Table 3.1(d) tax", "=B14"),
        (18, "GSTR-3B Table 4A(3) ITC", "=B14"),
    ]
    for row, label, val in rows:
        style(ws, f"A{row}", value=label, bold=True, fill=GREY, border=True)
        cell = style(ws, f"B{row}", value=val, border=True)
        if row in (10, 12, 13, 14, 15, 16, 17, 18):
            cell.number_format = money_fmt
        if row == 11:
            cell.number_format = pct_fmt
        ws.merge_cells(f"B{row}:C{row}")

    ws.merge_cells("A20:C22")
    style(ws, "A20",
          value=("Do not add ₹925 to the GTA NEFT. Pay them ₹18,500. Pay ₹925 from the "
                 "electronic cash ledger. Keep LR + SI-014 + PV-014 + UTR in one folder. "
                 "Copy these numbers into the free self invoice and payment voucher files."),
          border=True, wrap=True, italic=True)
    ws.row_dimensions[20].height = 22
    ws.row_dimensions[21].height = 22
    ws.row_dimensions[22].height = 22
    wb.save(path)


def write_register(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "RCM Register"
    ws.sheet_view.showGridLines = False
    headers = ["Date", "Self Invoice No.", "Payment Voucher No.", "Supplier",
               "Reason (GTA / URD / Advocate)", "Taxable (₹)", "GST %",
               "GST to govt (₹)", "Paid to supplier (₹)", "UTR", "GSTR-3B month"]
    widths = [12, 16, 18, 24, 26, 14, 10, 16, 18, 16, 14]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.merge_cells("A1:K1")
    style(ws, "A1", value="RCM DOCUMENT REGISTER — join SI and PV for the year",
          bold=True, size=16, color=WHITE, fill=BRAND, align="center")
    ws.row_dimensions[1].height = 28
    for i, h in enumerate(headers, start=1):
        style(ws, f"{get_column_letter(i)}3", value=h, bold=True, color=WHITE,
              fill=BRAND, align="center", border=True, wrap=True)
    ws.row_dimensions[3].height = 30
    # Sample row matching the GTA example
    style(ws, "A4", value="15-09-2026", border=True)
    style(ws, "B4", value="SI-014", border=True)
    style(ws, "C4", value="PV-014", border=True)
    style(ws, "D4", value="Sample Roadways", border=True)
    style(ws, "E4", value="GTA", border=True)
    style(ws, "F4", value=18500, border=True, number_format=money_fmt)
    style(ws, "G4", value=0.05, border=True, number_format=pct_fmt)
    style(ws, "H4", value="=IF(F4=\"\",\"\",ROUND(F4*G4,2))", border=True, number_format=money_fmt)
    style(ws, "I4", value="=F4", border=True, number_format=money_fmt)
    style(ws, "J4", value="[UTR]", border=True)
    style(ws, "K4", value="Sep-2026", border=True)
    for row in range(5, 25):
        for col in range(1, 12):
            style(ws, f"{get_column_letter(col)}{row}", border=True)
        style(ws, f"H{row}",
              value=f'=IF(F{row}="","",ROUND(F{row}*G{row},2))',
              number_format=money_fmt, border=True)
        style(ws, f"I{row}", value=f'=IF(F{row}="","",F{row})',
              number_format=money_fmt, border=True)
    ws.merge_cells("A26:K26")
    style(ws, "A26",
          value="One row per RCM inward supply. Never reuse SI or PV numbers. Totals are for your GSTR-3B check, not a substitute for the return.",
          italic=True, size=10)
    wb.save(path)


GUIDE = """TemplateHub GST Compliance Pack
================================
One-time purchase. Free single files stay free on the site.

WHAT IS IN THIS ZIP
1. 01-gst-invoice-template.xlsx
   Intra-state tax invoice (CGST + SGST). Same file as the free download.
2. 02-self-invoice-rcm-template.xlsx
   RCM self invoice (Section 31(3)(f)). Same as the free download.
3. 03-rcm-payment-voucher-template.xlsx
   Payment voucher (Section 31(3)(g)). Same as the free download.
4. 04-gst-invoice-igst-template.xlsx  (PACK ONLY)
   Inter-state tax invoice. One IGST column at the full rate. Use when
   place of supply is another state.
5. 05-gta-freight-worked-example.xlsx  (PACK ONLY)
   Filled GTA freight of Rs 18,500 at 5%. Shows what you pay the GTA
   vs what you pay the government, and the GSTR-3B tables.
6. 06-rcm-document-register.xlsx  (PACK ONLY)
   Year log that joins Self Invoice No. and Payment Voucher No.

HOW TO USE (RCM)
- Receive goods/services from an unregistered supplier → fill file 02, new SI-xxx.
- Pay the supplier → fill file 03, new PV-xxx, copy the same SI number.
- Log both numbers on file 06 the same day.
- Pay RCM GST from the electronic cash ledger. Do not add GST to the supplier NEFT.

This pack is a spreadsheet kit, not legal advice. Confirm rates (especially GTA
5% vs 12%) with your CA.

Questions: hello@templatehub.co.in
"""


def main() -> None:
    PACK_DIR.mkdir(parents=True, exist_ok=True)
    igst = PACK_DIR / "gst-invoice-igst-template.xlsx"
    gta = PACK_DIR / "gta-freight-worked-example.xlsx"
    reg = PACK_DIR / "rcm-document-register.xlsx"
    write_igst_invoice(igst)
    write_gta_example(gta)
    write_register(reg)

    zip_path = PACK_DIR / ZIP_NAME
    files = [
        (ROOT / "public/files/gst-invoice-template.xlsx", "01-gst-invoice-template.xlsx"),
        (ROOT / "public/files/self-invoice-rcm-template.xlsx", "02-self-invoice-rcm-template.xlsx"),
        (ROOT / "public/files/rcm-payment-voucher-template.xlsx", "03-rcm-payment-voucher-template.xlsx"),
        (igst, "04-gst-invoice-igst-template.xlsx"),
        (gta, "05-gta-freight-worked-example.xlsx"),
        (reg, "06-rcm-document-register.xlsx"),
    ]
    missing = [str(src) for src, _ in files if not src.exists()]
    if missing:
        raise SystemExit("Missing source files:\n" + "\n".join(missing))

    with ZipFile(zip_path, "w", ZIP_DEFLATED) as zf:
        zf.writestr("00-START-HERE.txt", GUIDE)
        for src, name in files:
            zf.write(src, name)
    print(f"Saved: {zip_path} ({zip_path.stat().st_size} bytes)")


if __name__ == "__main__":
    main()

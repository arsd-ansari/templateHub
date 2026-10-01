"""
Generates an RCM Bill format (.xlsx) — the bill a registered supplier issues
when GST is payable by the recipient (for example a registered GTA that has
not opted for forward charge).

This is not a self invoice. The supplier issues and signs it. GST is shown
so the recipient can pay the government, and it is not added to the amount
the supplier collects.

Run:  python3 scripts/generate_rcm_bill.py
Output: public/files/rcm-bill-format.xlsx
"""

from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

BRAND = "0F766E"
BRAND_LIGHT = "D7EEEB"
GREY = "F1F5F9"
WHITE = "FFFFFF"
DARK = "16181D"
ACCENT = "B45309"

thin = Side(style="thin", color="C9D2DC")
box = Border(left=thin, right=thin, top=thin, bottom=thin)

money_fmt = "#,##0.00"
pct_fmt = "0%"

wb = Workbook()
ws = wb.active
ws.title = "RCM Bill"
ws.sheet_view.showGridLines = False
ws.page_setup.orientation = "landscape"
ws.page_setup.fitToPage = True
ws.page_setup.fitToWidth = 1
ws.page_setup.fitToHeight = 1
ws.page_setup.paperSize = ws.PAPERSIZE_A4
ws.sheet_properties.pageSetUpPr.fitToPage = True
ws.print_title_rows = "1:2"
ws.page_setup.horizontalCentered = True
ws.oddHeader.left.text = "RCM Bill — tax payable by recipient"
ws.oddFooter.right.text = "Page &P of &N"
ws.page_margins.left = 0.4
ws.page_margins.right = 0.4
ws.page_margins.top = 0.6
ws.page_margins.bottom = 0.5

widths = [6, 36, 12, 8, 14, 16, 8, 14, 14, 18]
for i, w in enumerate(widths, start=1):
    ws.column_dimensions[get_column_letter(i)].width = w


def style_cell(ref, *, value=None, bold=False, size=11, color=DARK,
               fill=None, align="left", border=False, wrap=False,
               number_format=None, italic=False):
    c = ws[ref]
    if value is not None:
        c.value = value
    c.font = Font(name="Calibri", bold=bold, size=size, color=color, italic=italic)
    c.alignment = Alignment(horizontal=align, vertical="center", wrap_text=wrap)
    if fill:
        c.fill = PatternFill("solid", fgColor=fill)
    if border:
        c.border = box
    if number_format:
        c.number_format = number_format
    return c


ws.merge_cells("A1:J1")
style_cell("A1", value="RCM BILL  ·  TAX PAYABLE BY RECIPIENT", bold=True, size=20,
           color=WHITE, fill=BRAND, align="center")
ws.row_dimensions[1].height = 34

ws.merge_cells("A2:J2")
style_cell(
    "A2",
    value=("Issued by a registered supplier. Rule 46 requires this bill to say whether tax is payable on reverse charge. "
           "Do not collect GST — the recipient pays it to the government."),
    italic=True, size=10, color=ACCENT, align="center", wrap=True
)
ws.row_dimensions[2].height = 32

ws.merge_cells("A4:E4")
style_cell("A4", value="SUPPLIER (Registered — issuer of this bill)",
           bold=True, color=WHITE, fill=BRAND, border=True)
for row, key in ((5, "Legal Name"), (6, "Address"), (7, "GSTIN"), (8, "State & Code")):
    ws.merge_cells(f"A{row}:E{row}")
    style_cell(f"A{row}", value=f"{key}: [ ]", border=True, wrap=True)

meta = [
    ("Bill No.", "[RCM-001]"),
    ("Bill Date", "[DD-MM-YYYY]"),
    ("Consignment / LR No.", "[LR-1001]"),
    ("Vehicle No.", "[MH01AB1234]"),
    ("Place of Supply", "[Recipient's State (Code)]"),
    ("Reverse Charge", "Yes"),
]
r = 4
for label, val in meta:
    style_cell(f"G{r}", value=label, bold=True, fill=GREY, border=True)
    ws.merge_cells(f"H{r}:J{r}")
    style_cell(
        f"H{r}", value=val, border=True,
        bold=(label == "Reverse Charge"),
        color=(ACCENT if label == "Reverse Charge" else DARK)
    )
    r += 1

ws.merge_cells("A10:E10")
style_cell("A10", value="RECIPIENT (pays GST to the government)",
           bold=True, color=WHITE, fill=BRAND, border=True)
for row, key in ((11, "Legal Name"), (12, "Address"), (13, "GSTIN"), (14, "State & Code")):
    ws.merge_cells(f"A{row}:E{row}")
    style_cell(f"A{row}", value=f"{key}: [ ]", border=True, wrap=True)

ws.merge_cells("F10:J10")
style_cell("F10", value="MOVEMENT", bold=True, color=WHITE, fill=BRAND, border=True)
for row, key in ((11, "From (origin)"), (12, "To (destination)")):
    ws.merge_cells(f"F{row}:J{row}")
    style_cell(f"F{row}", value=f"{key}: [ ]", border=True, wrap=True)
ws.merge_cells("F13:J14")
style_cell(
    "F13",
    value=("Sample below is intra-state GTA freight: taxable ₹18,500 at 5% GST "
           "= CGST ₹462.50 + SGST ₹462.50. That ₹925 is not added to the freight you collect."),
    border=True, wrap=True, italic=True, size=10, color=ACCENT
)
ws.row_dimensions[13].height = 22
ws.row_dimensions[14].height = 32

HEADER_ROW = 16
headers = [
    "S.No", "Description of Goods/Services", "SAC", "Qty", "Rate",
    "Taxable Value", "GST %", "CGST", "SGST", "Payable to supplier"
]
for i, h in enumerate(headers, start=1):
    col = get_column_letter(i)
    style_cell(f"{col}{HEADER_ROW}", value=h, bold=True, color=WHITE,
               fill=BRAND, align="center", border=True, wrap=True)
ws.row_dimensions[HEADER_ROW].height = 30

FIRST = HEADER_ROW + 1
N_ROWS = 6
LAST = FIRST + N_ROWS - 1

# One worked line. Mumbai → Pune is the same state, so CGST + SGST.
samples = [
    (1, "Transportation of goods by road (GTA) — Mumbai to Pune", "996511", 1, 18500, 0.05),
]

for idx, row in enumerate(range(FIRST, LAST + 1)):
    sample = samples[idx] if idx < len(samples) else None
    style_cell(f"A{row}", value=(sample[0] if sample else None), align="center", border=True)
    style_cell(f"B{row}", value=(sample[1] if sample else None), border=True, wrap=True)
    style_cell(f"C{row}", value=(sample[2] if sample else None), align="center", border=True)
    style_cell(f"D{row}", value=(sample[3] if sample else None), align="center", border=True)
    style_cell(f"E{row}", value=(sample[4] if sample else None), align="right", border=True, number_format=money_fmt)
    style_cell(f"G{row}", value=(sample[5] if sample else None), align="center", border=True, number_format=pct_fmt)
    style_cell(
        f"F{row}",
        value=f'=IF(OR(D{row}="",E{row}=""),"",ROUND(D{row}*E{row},2))',
        align="right", border=True, number_format=money_fmt
    )
    style_cell(
        f"H{row}",
        value=f'=IF(F{row}="","",ROUND(F{row}*G{row}/2,2))',
        align="right", border=True, number_format=money_fmt
    )
    style_cell(
        f"I{row}",
        value=f'=IF(F{row}="","",ROUND(F{row}*G{row}/2,2))',
        align="right", border=True, number_format=money_fmt
    )
    # Supplier collects taxable value only. GST stays in H and I for the recipient.
    style_cell(
        f"J{row}",
        value=f'=IF(F{row}="","",F{row})',
        align="right", border=True, number_format=money_fmt
    )
    if idx % 2 == 1:
        for i in range(1, 11):
            ws[f"{get_column_letter(i)}{row}"].fill = PatternFill("solid", fgColor=GREY)
    ws.row_dimensions[row].height = 20
ws.row_dimensions[FIRST].height = 32


def total_row(row, label, formula, *, strong=False, accent=False):
    ws.merge_cells(f"A{row}:G{row}")
    fill = BRAND if strong else ("FEF3C7" if accent else BRAND_LIGHT)
    color = WHITE if strong else (ACCENT if accent else DARK)
    style_cell(f"A{row}", value=label, bold=True, align="right", fill=fill, color=color, border=True)
    ws.merge_cells(f"H{row}:I{row}")
    style_cell(f"H{row}", border=True, fill=(BRAND if strong else None))
    style_cell(
        f"J{row}", value=formula, bold=True, align="right", border=True, number_format=money_fmt,
        fill=fill, color=color, size=(12 if strong else 11)
    )


base = LAST + 1
total_row(base, "Amount payable to supplier (do not add GST)", f"=ROUND(SUM(J{FIRST}:J{LAST}),2)", strong=True)
total_row(base + 1, "CGST payable by recipient to government", f"=ROUND(SUM(H{FIRST}:H{LAST}),2)", accent=True)
total_row(base + 2, "SGST payable by recipient to government", f"=ROUND(SUM(I{FIRST}:I{LAST}),2)", accent=True)
total_row(base + 3, "Total GST under reverse charge (not collected on this bill)", f"=ROUND(J{base+1}+J{base+2},2)", accent=True)

words = base + 5
ws.merge_cells(f"A{words}:J{words}")
style_cell(
    f"A{words}",
    value="Amount payable to supplier, in words: [ Rupees ______________________ Only ]",
    bold=True, italic=True
)

decl = words + 2
ws.merge_cells(f"A{decl}:J{decl}")
style_cell(f"A{decl}", value="DECLARATION", bold=True, color=WHITE, fill=BRAND, border=True, align="center")
lines = [
    "Tax is payable on reverse charge basis by the recipient. The supplier has not collected GST on this bill.",
    "For an inter-state supply, do not split the tax: charge IGST at the full GST % on the taxable value, and leave the CGST and SGST columns unused.",
    "If the supplier is not registered, do not use this bill. The recipient issues a self invoice under Section 31(3)(f) instead.",
]
for i, line in enumerate(lines):
    ws.merge_cells(f"A{decl+1+i}:J{decl+1+i}")
    style_cell(f"A{decl+1+i}", value=line, border=True, size=10, wrap=True)
    ws.row_dimensions[decl + 1 + i].height = 28

sign = decl + len(lines) + 2
ws.merge_cells(f"G{sign}:J{sign}")
style_cell(f"G{sign}", value="For [Supplier Name]", bold=True, align="right")
ws.merge_cells(f"G{sign+2}:J{sign+2}")
style_cell(f"G{sign+2}", value="Authorised Signatory", align="right", color=BRAND)

ws.freeze_panes = f"A{FIRST}"
ws.print_area = f"A1:J{sign+2}"

out = Path("public/files/rcm-bill-format.xlsx")
out.parent.mkdir(parents=True, exist_ok=True)
wb.save(out)
print(f"Saved: {out.resolve()}")
